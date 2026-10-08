# -*- coding: utf-8 -*-
"""
Приймальні тести TENDERWIN LEAD & MAIL v1.0.

Мережі немає: стрічка Prozorro і Gmail підмінені. Перевіряються саме ті
речі, заради яких писався новий скрипт:

  * у розсилку потрапляють лише події вікна прогону;
  * стара подія не стає сьогоднішньою через документ, stand-still чи
    відхилення іншого учасника;
  * одна компанія — один лист, з урахуванням історії попереднього движка;
  * у тестовому режимі лист не може піти на адресу клієнта;
  * документи не завантажуються взагалі.

Запуск:  python3 test_lead_mail.py
"""
from __future__ import annotations

import inspect
import json
import os
import sqlite3
import sys
import tempfile
import unittest
from datetime import datetime, timedelta

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import tenderwin_lead_mail as M          # noqa: E402
import tenderwin_lead_engine as E        # noqa: E402

E.ALLOW_COLAB_SECRETS = False            # жодних справжніх секретів у тестах

KYIV = M.KYIV
TODAY = datetime(2026, 9, 22, 18, 0, tzinfo=KYIV)       # «зараз» у тестах


def iso(day: int, hour: int = 12, minute: int = 0, month: int = 9) -> str:
    return datetime(2026, month, day, hour, minute, tzinfo=KYIV).isoformat()


def supplier(code="12345678", name="ТОВ «БУДІНВЕСТ»", email="office@bud.example",
             person="Коваленко Іван Петрович", scheme="UA-EDR"):
    return {"name": name,
            "identifier": {"scheme": scheme, "id": code, "legalName": name},
            "address": {"region": "Київська", "locality": "Київ"},
            "contactPoint": {"name": person, "email": email,
                             "telephone": "+380671234567"}}


def bid(bid_id="b1", **kw):
    return {"id": bid_id, "tenderers": [supplier(**kw)],
            "documents": [{"id": "d1", "title": "Пропозиція.pdf",
                           "url": "https://example/d1"}],
            "lotValues": []}


def award(aid="a1", bid_id="b1", date=None, cp_start=None, cp_end=None,
          docs=None, status="unsuccessful", **kw):
    obj = {"id": aid, "bid_id": bid_id, "status": status,
           "date": date or iso(22, 11),
           "title": "Відхилення тендерної пропозиції",
           "description": "Учасник не надав довідку про працівників.",
           "documents": docs if docs is not None else
           [{"id": "p1", "title": "Протокол.pdf", "url": "https://example/p1",
             "datePublished": date or iso(22, 11)}],
           "suppliers": [supplier(**kw)]}
    if cp_start or cp_end:
        obj["complaintPeriod"] = {"startDate": cp_start or date or iso(22, 11),
                                  "endDate": cp_end or iso(27, 0)}
    return obj


def qualification(qid="q1", bid_id="b1", date=None, cp_start=None, cp_end=None):
    obj = {"id": qid, "bidID": bid_id, "status": "unsuccessful",
           "date": date or iso(18, 10),
           "title": "Відхилення",
           "description": "Не підтверджено кваліфікацію.",
           "documents": []}
    if cp_start or cp_end:
        obj["complaintPeriod"] = {"startDate": cp_start or iso(22, 0, 5),
                                  "endDate": cp_end or iso(27, 0)}
    return obj


def tender(uid="t1", ua="UA-2026-09-01-000001-a", awards=(), quals=(), bids=(),
           status="active.qualification", cpv="45453000-7", value=5_000_000,
           method="aboveThreshold"):
    return {"id": uid, "tenderID": ua, "status": status,
            "procurementMethodType": method,
            "title": "Капітальний ремонт покрівлі",
            "procuringEntity": {"name": "Управління освіти",
                                "identifier": {"id": "00000001"}},
            "items": [{"classification": {"id": cpv}}],
            "value": {"amount": value},
            "documents": [{"id": "td1", "title": "Тендерна документація.docx",
                           "url": "https://example/td1",
                           "datePublished": iso(1, 9)}],
            "bids": list(bids), "awards": list(awards),
            "qualifications": list(quals)}


class FakeApi(M.Prozorro):
    """Стрічка і закупівлі з памʼяті. Справжній код обходу при цьому працює."""

    def __init__(self, tenders, feed_pages_fail_at=None, broken=()):
        self.store = {t["id"]: t for t in tenders}
        self.feed_fail = feed_pages_fail_at
        self.broken = set(broken)
        self.session = None
        self.api = "http://fake/tenders"
        self.verbose = False
        self.page = 0

    def _get(self, url, params=None):
        if params is not None:                      # сторінка стрічки
            self.page += 1
            if self.feed_fail and self.page >= self.feed_fail:
                return None, "HTTP_503"
            if self.page > 1:
                # Справжня стрічка закінчується порожньою сторінкою.
                return {"data": [], "next_page": {"offset": "кінець"}}, ""
            rows = [{"id": t["id"], "status": t["status"],
                     "procurementMethodType": t["procurementMethodType"]}
                    for t in self.store.values()]
            return {"data": rows, "next_page": {"offset": "x"}}, ""
        uid = url.rsplit("/", 1)[1]
        if uid in self.broken:
            return None, "HTTP_404"
        return {"data": self.store[uid]}, ""


class FakeResult:
    def __init__(self, ok=True, unknown=False, error=""):
        self.ok, self.unknown, self.error = ok, unknown, error
        self.message_id = "gmail-1" if ok else None
        self.thread_id = "thr-1" if ok else None
        self.provider_timestamp = M.now_iso()


class FakeTransport:
    name = "FAKE"

    def __init__(self, result=None, labels=("SENT",), raise_on_find=False):
        self.result = result or FakeResult()
        self.labels = labels
        self.sent = []
        self.raise_on_find = raise_on_find
        self._read_api = self                    # verify_sent зайде сюди

    # --- транспорт ---
    def send(self, raw, message_id_header):
        self.sent.append((message_id_header, raw))
        return self.result

    def find_by_message_id(self, message_id_header):
        if self.raise_on_find:
            raise RuntimeError("немає доступу")
        return ({"id": "gmail-1", "threadId": "thr-1"}
                if any(message_id_header == m for m, _ in self.sent) else None)

    # --- мінімальний макет Gmail API для перевірки мітки SENT ---
    def users(self):
        return self

    def messages(self):
        return self

    def get(self, **kwargs):
        labels = self.labels
        class _Exec:
            def execute(self_inner):
                return {"labelIds": list(labels)}
        return _Exec()


class Base(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="lm_")
        self.db_path = os.path.join(self.tmp, "state.db")
        self.saved = {k: getattr(M, k) for k in (
            "MODE", "LETTER_VERSION", "TEST_RECIPIENTS", "MIN_LOT_VALUE",
            "SKIP_WITHOUT_COMPLAINT_ROUTE", "LETTER_FORMAT",
            "WINDOW_MODE", "SHOW_COMPLAINT_DEADLINE", "ALLOW_MEDIUM_RELIABILITY",
            "SHOW_OPT_OUT_LINE", "DB_PATH", "OUT_DIR", "LEGACY_DB_PATH")}
        M.TEST_RECIPIENTS = ["test1@example.com", "test2@example.com"]
        M.DB_PATH, M.OUT_DIR = self.db_path, os.path.join(self.tmp, "vyhid")
        M.LEGACY_DB_PATH = os.path.join(self.tmp, "немає.db")
        # Годинник зафіксовано: інакше строки оскарження з фікстур (27.09)
        # «минають» у реальному часі, і тести падають щодня після них.
        self._real_now = M.now
        M.now = lambda: TODAY

    def tearDown(self):
        for key, value in self.saved.items():
            setattr(M, key, value)
        M.now = self._real_now

    def engine(self):
        eng = M.LeadMail(db_path=self.db_path,
                         out_dir=os.path.join(self.tmp, "vyhid"), verbose=False)
        eng.prepare()
        return eng

    def run_scan(self, tenders, *, day_now=TODAY, window=None, api=None, **kw):
        eng = self.engine()
        win = window or M.Window(*M.day_bounds(day_now), mode="TODAY_ONLY")
        win = M.Window(win.start, min(win.end, day_now), win.mode, win.tail_start)
        api = api or FakeApi(tenders, **kw)
        result = api.scan(win)
        eng.repo.start_run(eng.run_id, win)
        eng.ingest(result)
        return eng, result


# ---------------------------------------------------------------------------
class TestEventTime(Base):
    """Час події: одне поле, джерело, обґрунтування, надійність."""

    def test_award_two_fields_agree_is_high(self):
        e = M.event_time(award(date=iso(22, 11), cp_start=iso(22, 11)), "awards", TODAY)
        self.assertEqual(e.reliability, "HIGH")
        self.assertIn("award.date", e.source)

    def test_award_fields_disagree_is_conflict(self):
        e = M.event_time(award(date=iso(19, 16), cp_start=iso(22, 9)), "awards", TODAY)
        self.assertEqual(e.reliability, "CONFLICT")

    def test_award_single_field_is_medium(self):
        e = M.event_time({"date": iso(22, 11)}, "awards", TODAY)
        self.assertEqual(e.reliability, "MEDIUM")

    def test_qualification_never_uses_complaint_period(self):
        e = M.event_time(qualification(date=iso(18, 10), cp_start=iso(22, 0)),
                         "qualifications", TODAY)
        self.assertEqual(e.time.day, 18, "stand-still не є часом рішення")
        self.assertEqual(e.reliability, "MEDIUM")

    def test_no_date_is_undetermined(self):
        e = M.event_time({"status": "unsuccessful"}, "awards", TODAY)
        self.assertEqual(e.reliability, "UNDETERMINED")

    def test_future_time_is_undetermined(self):
        e = M.event_time({"date": iso(23, 12)}, "awards", TODAY)
        self.assertEqual(e.reliability, "UNDETERMINED")

    def test_naive_datetime_is_rejected(self):
        self.assertIsNone(M.parse_dt("2026-09-22T11:00:00"))
        e = M.event_time({"date": "2026-09-22T11:00:00"}, "awards", TODAY)
        self.assertEqual(e.reliability, "UNDETERMINED")

    def test_document_dates_are_never_the_event_time(self):
        obj = award(date=iso(19, 16), cp_start=iso(19, 16),
                    docs=[{"id": "p", "title": "Протокол.pdf",
                           "url": "u", "datePublished": iso(22, 9)}])
        e = M.event_time(obj, "awards", TODAY)
        self.assertEqual(e.time.day, 19)


class TestWindow(Base):
    def test_day_bounds_are_local_midnights(self):
        start, end = M.day_bounds(TODAY)
        self.assertEqual((start.hour, start.minute), (0, 0))
        self.assertEqual(start.date().isoformat(), "2026-09-22")
        self.assertEqual(end.date().isoformat(), "2026-09-23")

    def test_dst_day_has_twenty_five_hours(self):
        start, end = M.day_bounds(datetime(2026, 10, 25, 12, tzinfo=KYIV))
        self.assertEqual(M.span_hours(start, end), 25.0)
        self.assertEqual((start.isoformat(), end.isoformat()),
                         ("2026-10-25T00:00:00+03:00", "2026-10-26T00:00:00+02:00"))

    def test_spring_day_has_twenty_three_hours(self):
        start, end = M.day_bounds(datetime(2027, 3, 28, 12, tzinfo=KYIV))
        self.assertEqual(M.span_hours(start, end), 23.0)
        self.assertEqual((start.isoformat(), end.isoformat()),
                         ("2027-03-28T00:00:00+02:00", "2027-03-29T00:00:00+03:00"))

    def test_window_report_line_names_the_time_shift(self):
        win = M.Window(*M.day_bounds(datetime(2026, 10, 25, 12, tzinfo=KYIV)),
                       mode="TODAY_ONLY")
        self.assertIn("доба 25 год", win.human())
        self.assertNotIn("год —", M.Window(*M.day_bounds(TODAY),
                                           mode="TODAY_ONLY").human())

    def test_run_ids_do_not_collide_within_one_second(self):
        ids = {M.LeadMail(db_path=self.db_path,
                          out_dir=os.path.join(self.tmp, "vyhid"),
                          verbose=False).run_id for _ in range(5)}
        self.assertEqual(len(ids), 5)

    def test_boundaries_include_midnight_and_exclude_previous_second(self):
        win = M.Window(*M.day_bounds(TODAY), mode="TODAY_ONLY")
        self.assertTrue(win.contains(datetime(2026, 9, 22, 0, 0, tzinfo=KYIV)))
        self.assertFalse(win.contains(datetime(2026, 9, 21, 23, 59, 59, tzinfo=KYIV)))
        self.assertTrue(win.contains(datetime(2026, 9, 22, 23, 59, 59, tzinfo=KYIV)))
        self.assertFalse(win.contains(datetime(2026, 9, 23, 0, 0, tzinfo=KYIV)))

    def test_tail_of_previous_day_is_counted_separately(self):
        win = M.Window(*M.day_bounds(TODAY), mode="TODAY_ONLY",
                       tail_start=datetime(2026, 9, 21, 18, tzinfo=KYIV))
        self.assertTrue(win.in_tail(datetime(2026, 9, 21, 20, tzinfo=KYIV)))
        self.assertFalse(win.in_tail(datetime(2026, 9, 20, 20, tzinfo=KYIV)))


class TestSelection(Base):
    """Головне: у роботу йдуть лише сьогоднішні події."""

    def test_old_and_new_rejection_in_one_tender(self):
        t = tender(bids=[bid("b1", code="11111111"), bid("b2", code="22222222")],
                   awards=[award("a-old", "b1", date=iso(10, 12), cp_start=iso(10, 12),
                                 code="11111111"),
                           award("a-new", "b2", date=iso(22, 11), cp_start=iso(22, 11),
                                 code="22222222")])
        eng, result = self.run_scan([t])
        codes = sorted(f.edrpou for f in result.events)
        eng.close()
        self.assertEqual(codes, ["22222222"], "старе відхилення не є сьогоднішнім")

    def test_reuploaded_protocol_does_not_revive_old_rejection(self):
        t = tender(uid="t2", ua="UA-2026-09-05-000002-a",
                   bids=[bid("b3", code="33333333")],
                   awards=[award("a-re", "b3", date=iso(19, 16), cp_start=iso(19, 16),
                                 code="33333333",
                                 docs=[{"id": "p", "title": "Рішення.pdf",
                                        "url": "u", "datePublished": iso(22, 9)}])])
        eng, result = self.run_scan([t])
        eng.close()
        self.assertEqual(result.events, [])
        self.assertEqual(result.counters["старіші"], 1)

    def test_standstill_today_does_not_revive_old_qualification(self):
        t = tender(uid="t3", ua="UA-2026-09-02-000003-a", method="aboveThresholdEU",
                   status="active.pre-qualification.stand-still",
                   bids=[bid("b4", code="44444444")],
                   quals=[qualification("q1", "b4", date=iso(18, 10),
                                        cp_start=iso(22, 0, 5))])
        eng, result = self.run_scan([t])
        eng.close()
        self.assertEqual(result.events, [])

    def test_todays_rejection_with_old_document_is_taken(self):
        t = tender(uid="t4", ua="UA-2026-09-03-000004-a",
                   bids=[bid("b5", code="55555555")],
                   awards=[award("a-min", "b5", date=iso(22, 14), cp_start=iso(22, 14),
                                 code="55555555",
                                 docs=[{"id": "n24", "title": "Вимога 24 год.pdf",
                                        "url": "u", "datePublished": iso(15, 9)},
                                       {"id": "p5", "title": "Протокол.pdf",
                                        "url": "u2", "datePublished": iso(22, 14)}])])
        eng, result = self.run_scan([t])
        eng.close()
        self.assertEqual(len(result.events), 1)
        self.assertEqual(result.events[0].event_time.day, 22)
        self.assertEqual(result.events[0].time_reliability, "HIGH")

    def test_conflicting_dates_go_to_review_not_to_letters(self):
        t = tender(bids=[bid("b1")],
                   awards=[award("a1", "b1", date=iso(22, 11), cp_start=iso(19, 9))])
        eng, result = self.run_scan([t])
        letters = eng.build_letters()
        states = {r["state"] for r in eng.db.q("SELECT state FROM events")}
        eng.close()
        self.assertEqual(result.events, [])
        self.assertEqual(letters["листів"], 0)
        self.assertEqual(states, {"REVIEW"})

    def test_all_industries_and_small_value_are_included(self):
        t = tender(uid="med", ua="UA-2026-09-22-000010-a", cpv="33141000-0",
                   value=250_000, bids=[bid("b1", code="77777777")],
                   awards=[award("a1", "b1", date=iso(22, 10), cp_start=iso(22, 10),
                                 code="77777777")])
        eng, result = self.run_scan([t])
        eng.close()
        self.assertEqual(len(result.events), 1, "фільтра за галуззю й сумою немає")

    def test_below_threshold_is_not_even_requested(self):
        """Допорогову закупівлю скрипт не питає в Prozorro взагалі."""
        t = tender(uid="low", ua="UA-2026-09-23-000020-a", method="belowThreshold",
                   value=48_000, bids=[bid("b1", code="66666666")],
                   awards=[award("a1", "b1", date=iso(22, 10), cp_start=iso(22, 10),
                                 code="66666666")])
        api = FakeApi([t])
        eng, result = self.run_scan([t], api=api)
        eng.close()
        self.assertEqual(result.events, [])
        self.assertEqual(result.counters["пропущено_за_процедурою"], 1)
        self.assertEqual(result.skipped_types, {"belowThreshold": 1})
        self.assertEqual(result.counters["закупівель_перевірено"], 0,
                         "закупівлю навіть не завантажували")

    def test_simplified_and_price_quotation_are_skipped(self):
        сторінки = [tender(uid="pq", ua="UA-2026-09-23-000021-a",
                           method="priceQuotation",
                           bids=[bid("b1", code="55555555")],
                           awards=[award("a1", "b1", date=iso(22, 10),
                                         cp_start=iso(22, 10), code="55555555")]),
                    tender(uid="sd", ua="UA-2026-09-23-000022-a",
                           method="simple.defense",
                           bids=[bid("b2", code="44444444")],
                           awards=[award("a2", "b2", date=iso(22, 10),
                                         cp_start=iso(22, 10), code="44444444")])]
        eng, result = self.run_scan(сторінки)
        eng.close()
        self.assertEqual(result.events, [])
        self.assertEqual(sorted(result.skipped_types), ["priceQuotation",
                                                        "simple.defense"])

    def test_unknown_new_procedure_is_skipped_and_named(self):
        """Новий невідомий тип не потрапляє в розсилку мовчки."""
        t = tender(uid="new", ua="UA-2026-09-23-000023-a", method="brandNewType2027",
                   bids=[bid("b1", code="33333333")],
                   awards=[award("a1", "b1", date=iso(22, 10), cp_start=iso(22, 10),
                                 code="33333333")])
        eng, result = self.run_scan([t])
        eng.close()
        self.assertEqual(result.events, [])
        self.assertEqual(result.skipped_types, {"brandNewType2027": 1})

    def test_open_tender_still_goes_through(self):
        for метод in ("aboveThreshold", "aboveThresholdUA", "aboveThresholdEU",
                      "negotiation"):
            with self.subTest(метод=метод):
                t = tender(uid="ok", ua="UA-2026-09-23-000024-a", method=метод,
                           bids=[bid("b1", code="22222222")],
                           awards=[award("a1", "b1", date=iso(22, 10),
                                         cp_start=iso(22, 10), code="22222222")])
                eng, result = self.run_scan([t])
                eng.close()
                self.assertEqual(len(result.events), 1, метод)

    def test_gap_retry_cannot_smuggle_a_below_threshold_tender(self):
        """Закупівля з прогалини теж перевіряється за типом процедури."""
        t = tender(uid="low2", ua="UA-2026-09-23-000025-a", method="belowThreshold",
                   bids=[bid("b1", code="11111199")],
                   awards=[award("a1", "b1", date=iso(22, 10), cp_start=iso(22, 10),
                                 code="11111199")])
        api = FakeApi([t])
        result = api.scan(M.Window(*M.day_bounds(TODAY), mode="TODAY_ONLY"),
                          retry_uids=["low2"])
        self.assertEqual(result.events, [])
        self.assertEqual(result.counters["пропущено_за_процедурою"], 1)

    def test_closed_tender_statuses_are_included(self):
        t = tender(uid="done", ua="UA-2026-09-22-000011-a", status="unsuccessful",
                   bids=[bid("b1", code="88888888")],
                   awards=[award("a1", "b1", date=iso(22, 10), cp_start=iso(22, 10),
                                 code="88888888")])
        eng, result = self.run_scan([t])
        eng.close()
        self.assertEqual(len(result.events), 1)

    def test_pre_award_statuses_are_skipped_in_feed(self):
        t = tender(uid="early", ua="UA-2026-09-22-000012-a", status="active.tendering",
                   bids=[bid("b1")], awards=[award("a1", "b1", date=iso(22, 10),
                                                   cp_start=iso(22, 10))])
        eng, result = self.run_scan([t])
        eng.close()
        self.assertEqual(result.counters["закупівель_перевірено"], 0)


class TestIdentity(Base):
    def test_supplier_and_tenderer_codes_must_match(self):
        t = tender(bids=[bid("b1", code="11111111")],
                   awards=[award("a1", "b1", date=iso(22, 11), cp_start=iso(22, 11),
                                 code="99999999")])
        eng, result = self.run_scan([t])
        eng.close()
        self.assertEqual(result.events, [])
        self.assertTrue(any("IDENTITY_CONFLICT" in (why or "")
                            for _, why in result.review))

    def test_joint_bid_is_recorded(self):
        b = bid("b1", code="11111111")
        b["tenderers"].append(supplier(code="22222222", name="ТОВ «Друга»"))
        t = tender(bids=[b], awards=[award("a1", "b1", date=iso(22, 11),
                                           cp_start=iso(22, 11), code="11111111")])
        eng, result = self.run_scan([t])
        eng.close()
        self.assertTrue(result.events[0].joint_bid)
        self.assertIn("спільна пропозиція", result.events[0].identity_note)

    def test_unknown_scheme_is_noted_but_not_fatal(self):
        t = tender(bids=[bid("b1", code="11111111", scheme="XM-DAC")],
                   awards=[award("a1", "b1", date=iso(22, 11), cp_start=iso(22, 11),
                                 code="11111111", scheme="XM-DAC")])
        eng, result = self.run_scan([t])
        eng.close()
        self.assertEqual(len(result.events), 1)
        self.assertIn("схема", result.events[0].identity_note)


class TestLetters(Base):
    def base_tender(self, code="11111111", uid="t1", ua="UA-2026-09-01-000001-a",
                    value=5_000_000):
        return tender(uid=uid, ua=ua, value=value, bids=[bid("b1", code=code)],
                      awards=[award("a1", "b1", date=iso(22, 11), cp_start=iso(22, 11),
                                    cp_end=iso(27, 0), code=code)])

    def test_one_company_one_letter_with_both_procurements(self):
        t1 = self.base_tender(uid="t1", ua="UA-2026-09-01-000001-a")
        t2 = self.base_tender(uid="t2", ua="UA-2026-09-02-000002-a", value=9_000_000)
        eng, _ = self.run_scan([t1, t2])
        stats = eng.build_letters()
        row = eng.db.one("SELECT * FROM outreach")
        eng.close()
        self.assertEqual(stats["листів"], 1)
        self.assertIn("UA-2026-09-01-000001-a", row["body_rendered"])
        self.assertIn("UA-2026-09-02-000002-a", row["body_rendered"])
        self.assertIn("UA-2026-09-02-000002-a", row["subject_rendered"],
                      "у темі — найдорожча закупівля")

    def test_second_run_does_not_write_again(self):
        t = self.base_tender()
        eng, _ = self.run_scan([t])
        eng.build_letters()
        eng.close()
        eng2, _ = self.run_scan([t])
        stats = eng2.build_letters()
        total = eng2.db.one("SELECT COUNT(*) n FROM outreach")["n"]
        eng2.close()
        self.assertEqual(stats["листів"], 0)
        self.assertEqual(total, 1)

    def test_owner_template_is_rendered_verbatim(self):
        M.LETTER_VERSION = "V4_OWNER"
        eng, _ = self.run_scan([self.base_tender()])
        eng.build_letters()
        row = eng.db.one("SELECT * FROM outreach")
        eng.close()
        self.assertIn("Вартість — 3 499 грн", row["body_rendered"])
        self.assertIn("Щасливий Віталій", row["body_rendered"])
        self.assertTrue(row["subject_rendered"].startswith("Чи обґрунтовано"))
        self.assertIn("Добрий день, Іване!", row["body_rendered"])

    def test_single_rejection_letter_matches_the_approved_text(self):
        """Золотий текст: один лист про одне відхилення — слово в слово."""
        M.LETTER_VERSION = "V4_OWNER"
        eng, _ = self.run_scan([self.base_tender()])
        eng.build_letters()
        row = eng.db.one("SELECT * FROM outreach")
        eng.close()
        очікувано = (
            "Добрий день, Іване!\n"
            "\n"
            "Звернув увагу, що вашу пропозицію в закупівлі "
            "UA-2026-09-01-000001-a «Капітальний ремонт покрівлі» було відхилено.\n"
            "\n"
            "Чи обґрунтоване рішення замовника та чи варто витрачати кошти на "
            "його оскарження?\n"
            "\n"
            "TenderWin пропонує аналіз вашого відхилення за 24 години. Ви "
            "отримаєте письмовий висновок: чи є підстави для оскарження, які "
            "існують ризики та що доцільно робити далі.\n"
            "\n"
            "Якщо оскарження недоцільне — прямо про це повідомимо.\n"
            "\n"
            "Вартість — 3 499 грн. Термін — після оплати й отримання "
            "необхідних документів.\n"
            "\n"
            "Якщо питання актуальне, просто відповідайте на цей лист.\n"
            "\n"
            "З повагою,\n"
            "Щасливий Віталій\n"
            "Радник з публічних закупівель TenderWin\n"
            "050 310 14 92\n")
        self.assertEqual(row["body_rendered"], очікувано)
        self.assertEqual(row["subject_rendered"],
                         "Чи обґрунтовано відхилили вашу пропозицію "
                         "UA-2026-09-01-000001-a?")

    def test_two_rejections_use_plural_forms(self):
        """Дві закупівлі — «ваші пропозиції в закупівлях», а не «пропозицію в закупівлі»."""
        M.LETTER_VERSION = "V4_OWNER"
        t1 = self.base_tender(uid="t1", ua="UA-2026-09-01-000001-a")
        t2 = self.base_tender(uid="t2", ua="UA-2026-09-02-000002-a", value=9_000_000)
        eng, _ = self.run_scan([t1, t2])
        eng.build_letters()
        row = eng.db.one("SELECT * FROM outreach")
        eng.close()
        self.assertIn("ваші пропозиції в закупівлях", row["body_rendered"])
        self.assertIn("аналіз ваших відхилень", row["body_rendered"])
        self.assertNotIn("вашу пропозицію в закупівлі", row["body_rendered"])

    def test_procurement_without_appeal_route_is_set_aside(self):
        """Без періоду оскарження лист не складається, лід лишається видимим."""
        t = tender(bids=[bid("b1")], awards=[award("a1", "b1", date=iso(22, 11))])
        eng, _ = self.run_scan([t])
        stats = eng.build_letters()
        стан = eng.db.one("SELECT state, state_reason FROM events")
        eng.close()
        self.assertEqual(stats["листів"], 0)
        self.assertEqual(stats["NO_COMPLAINT_ROUTE"], 1)
        self.assertEqual(стан["state"], "NO_COMPLAINT_ROUTE")

    def test_owner_can_include_them_with_one_setting(self):
        """Одне значення вмикає розсилку і на такі закупівлі."""
        M.SKIP_WITHOUT_COMPLAINT_ROUTE = False
        t = tender(bids=[bid("b1")], awards=[award("a1", "b1", date=iso(22, 11))])
        eng, _ = self.run_scan([t])
        stats = eng.build_letters()
        row = eng.db.one("SELECT * FROM outreach")
        eng.close()
        self.assertEqual(stats["листів"], 1)
        self.assertNotIn("оскарження можна до", row["body_rendered"],
                         "строку немає — рядок про строк не зʼявляється")

    def test_redteam_template_switches_by_one_value(self):
        M.LETTER_VERSION = "V5_REDTEAM"
        eng, _ = self.run_scan([self.base_tender()])
        eng.build_letters()
        row = eng.db.one("SELECT * FROM outreach")
        eng.close()
        self.assertEqual(row["template_version"], "V5_REDTEAM")
        self.assertIn("Побачив у Prozorro", row["body_rendered"])
        self.assertIn("подати скаргу на це рішення можна до", row["body_rendered"])

    def test_deadline_line_disappears_when_period_is_closed(self):
        M.LETTER_VERSION = "V5_REDTEAM"
        t = tender(bids=[bid("b1")],
                   awards=[award("a1", "b1", date=iso(22, 11), cp_start=iso(22, 11),
                                 cp_end=iso(20, 0))])
        eng, _ = self.run_scan([t])
        eng.build_letters()
        row = eng.db.one("SELECT * FROM outreach")
        eng.close()
        self.assertNotIn("подати скаргу", row["body_rendered"])

    def test_letter_never_contains_placeholders(self):
        eng, _ = self.run_scan([self.base_tender()])
        eng.build_letters()
        row = eng.db.one("SELECT * FROM outreach")
        eng.close()
        left = E._leftover_placeholders(row["subject_rendered"] + row["body_rendered"])
        self.assertEqual(left, [])

    def test_no_email_goes_to_review(self):
        b = bid("b1")
        b["tenderers"][0]["contactPoint"].pop("email")
        a = award("a1", "b1", date=iso(22, 11), cp_start=iso(22, 11))
        a["suppliers"][0]["contactPoint"].pop("email")
        eng, _ = self.run_scan([tender(bids=[b], awards=[a])])
        stats = eng.build_letters()
        reason = eng.db.one("SELECT state, state_reason FROM events")
        eng.close()
        self.assertEqual(stats["листів"], 0)
        self.assertEqual(reason["state"], "REVIEW")
        self.assertIn("e-mail", reason["state_reason"])

    def test_tender_without_complaint_period_gets_no_letter(self):
        t = tender(bids=[bid("b1")], awards=[award("a1", "b1", date=iso(22, 11))])
        eng, _ = self.run_scan([t])
        stats = eng.build_letters()
        state = eng.db.one("SELECT state FROM events")["state"]
        eng.close()
        self.assertEqual(stats["листів"], 0)
        self.assertEqual(state, "NO_COMPLAINT_ROUTE")

    def test_min_lot_value_filter_is_off_by_default_and_works_when_set(self):
        self.assertEqual(M.MIN_LOT_VALUE, 0)
        M.MIN_LOT_VALUE = 10_000_000
        eng, _ = self.run_scan([self.base_tender(value=5_000_000)])
        stats = eng.build_letters()
        eng.close()
        self.assertEqual(stats["листів"], 0)


class TestLetterV8(Base):
    """Лист V8 (23.09.2026): заморожена версія — текст не змінюється."""

    def setUp(self):
        super().setUp()
        M.LETTER_VERSION = "V8_OWNER"

    def лист(self, cp_end=None, title="Капітальний ремонт покрівлі"):
        t = tender(bids=[bid("b1")],
                   awards=[award("a1", "b1", date=iso(22, 11), cp_start=iso(22, 11),
                                 cp_end=cp_end or iso(27, 0))])
        t["title"] = title
        eng, _ = self.run_scan([t])
        eng.build_letters()
        row = eng.db.one("SELECT * FROM outreach")
        eng.close()
        return row

    def test_v8_still_renders_when_chosen(self):
        self.assertEqual(self.лист()["template_version"], "V8_OWNER")

    def test_previous_versions_stay_frozen(self):
        """Старі версії не редагуються: інакше не скажеш, що пішло раніше."""
        self.assertIn("Цього Вам буде достатньо",
                      M.ALL_TEMPLATES["V6_OWNER"].body)
        self.assertIn("Зазвичай цього достатньо",
                      M.ALL_TEMPLATES["V7_OWNER"].body)
        self.assertNotIn("Я уважно зіставлю", M.ALL_TEMPLATES["V7_OWNER"].body)
        self.assertIn("Я уважно зіставлю", M.ALL_TEMPLATES["V8_OWNER"].body)
        self.assertIn("Якщо питання актуальне — просто відповідайте на цей лист",
                      M.ALL_TEMPLATES["V8_OWNER"].body)

    def test_text_matches_the_approved_wording(self):
        """Золотий текст V8 — слово в слово, як затвердив власник."""
        row = self.лист()
        очікувано = (
            "Добрий день, Іване!\n"
            "\n"
            "Побачив у Prozorro, що пропозицію ТОВ «БУДІНВЕСТ» "
            "(ЄДРПОУ 12345678) в закупівлі UA-2026-09-01-000001-a "
            "«Капітальний ремонт покрівлі» відхилено "
            "(22 вересня 2026 року об 11:00).\n"
            "\n"
            "Питання після такого рішення зазвичай одне: чи справді підстави "
            "були і чи варто витрачати гроші на оскарження.\n"
            "\n"
            "Дам на нього дуже детальну письмову відповідь:\n"
            "\n"
            "• чи були підстави для відхилення,\n"
            "• які ризики,\n"
            "• що доцільно робити далі.\n"
            "\n"
            "Я уважно зіставлю те, що написав замовник у рішенні, з вимогами "
            "тендерної документації, вашими документами та практикою АМКУ. "
            "Якщо підстав для скарги немає, напишу про це прямо.\n"
            "\n"
            "Зазвичай цього достатньо, щоб самостійно подати скаргу в АМКУ "
            "(якщо це доцільно), не витрачаючи зайвих коштів на юридичний "
            "супровід.\n"
            "\n"
            "Вартість — 3 499 грн за одне рішення про відхилення, з усіма "
            "підставами, які в ньому названі.\n"
            "Висновок — протягом 24 годин після оплати.\n"
            "\n"
            "У Prozorro подати скаргу на це рішення можна до "
            "27 вересня 2026 року о 00:00.\n"
            "\n"
            "Якщо питання актуальне — просто відповідайте на цей лист.\n"
            "\n"
            "З повагою,\n"
            "Віталій Щасливий\n"
            "радник з публічних закупівель TenderWin\n"
            "+380 50 310 14 92 · tenderwin.in.ua\n")
        self.assertEqual(row["body_rendered"], очікувано)
        self.assertEqual(row["subject_rendered"],
                         "Відхилення у закупівлі UA-2026-09-01-000001-a: "
                         "чи є підстави для оскарження?")

    def test_html_has_the_same_words_as_the_text(self):
        """Оформлення не додає і не прибирає жодного речення."""
        import re as _re
        row = self.лист()
        html = row["body_html"]
        текст_з_html = _re.sub(r"<[^>]+>", " ", html)
        текст_з_html = _re.sub(r"\s+", " ", текст_з_html).replace("&amp;", "&")
        for речення in ("Питання після такого рішення зазвичай одне",
                        "Зазвичай цього достатньо, щоб самостійно подати скаргу",
                        "Вартість — 3 499 грн",
                        "протягом 24 годин після оплати",
                        "У Prozorro подати скаргу на це рішення можна до "
                        "27 вересня",
                        "Якщо питання актуальне — просто відповідайте на цей лист",
                        "Віталій Щасливий",
                        "+380 50 310 14 92"):
            self.assertIn(речення, текст_з_html, речення)

    def test_company_name_and_code_are_in_the_letter(self):
        row = self.лист()
        self.assertIn("ТОВ «БУДІНВЕСТ» (ЄДРПОУ 12345678)", row["body_rendered"])
        self.assertIn("ТОВ «БУДІНВЕСТ» (ЄДРПОУ 12345678)", row["body_html"])

    def test_individual_entrepreneur_gets_ipn_label(self):
        """ФОП має десятизначний код — це ІПН, а не ЄДРПОУ."""
        t = tender(bids=[bid("b1", code="1234567890", name="ФОП Сидоренко О. М.")],
                   awards=[award("a1", "b1", date=iso(22, 11), cp_start=iso(22, 11),
                                 cp_end=iso(27, 0), code="1234567890",
                                 name="ФОП Сидоренко О. М.")])
        eng, _ = self.run_scan([t])
        eng.build_letters()
        row = eng.db.one("SELECT * FROM outreach")
        eng.close()
        self.assertIn("(ІПН 1234567890)", row["body_rendered"])
        self.assertNotIn("ЄДРПОУ", row["body_rendered"])

    def test_exact_time_of_the_decision_is_stated(self):
        row = self.лист()
        self.assertIn("22 вересня 2026 року об 11:00", row["body_rendered"])
        self.assertIn("22 вересня 2026 року об 11:00", row["body_html"])

    def test_deadline_carries_date_and_time(self):
        row = self.лист(cp_end=iso(27, 14, 30))
        self.assertIn("можна до 27 вересня 2026 року о 14:30",
                      row["body_rendered"])
        self.assertIn("27 вересня 2026 року о 14:30", row["body_html"])

    def test_two_rejections_are_listed_each_with_its_own_time(self):
        """Два відхилення — два рядки: у кожного свій час рішення."""
        t1 = tender(uid="t1", ua="UA-2026-09-01-000001-a",
                    bids=[bid("b1")],
                    awards=[award("a1", "b1", date=iso(22, 11), cp_start=iso(22, 11),
                                  cp_end=iso(27, 0))])
        t2 = tender(uid="t2", ua="UA-2026-09-02-000002-a", value=9_000_000,
                    bids=[bid("b1")],
                    awards=[award("a2", "b1", date=iso(22, 15, 45),
                                  cp_start=iso(22, 15, 45), cp_end=iso(28, 0))])
        eng, _ = self.run_scan([t1, t2])
        eng.build_letters()
        row = eng.db.one("SELECT * FROM outreach")
        eng.close()
        тіло = row["body_rendered"]
        self.assertIn("що пропозиції ТОВ «БУДІНВЕСТ» (ЄДРПОУ 12345678) "
                      "відхилено:", тіло)
        self.assertIn("• UA-2026-09-01-000001-a", тіло)
        self.assertIn("22 вересня 2026 року об 11:00", тіло)
        self.assertIn("22 вересня 2026 року о 15:45", тіло)
        self.assertIn("найраніший строк", тіло)
        self.assertIn("<li", row["body_html"])

    def test_letter_is_not_built_without_the_time_of_the_decision(self):
        """Без часу рішення лист не складається — він назвав би невстановлене."""
        події = [M.Rejection(ua_id="UA-1", tender_id="t", tender_title="Ремонт",
                             buyer_name="", buyer_edrpou="", cpv="", method_type="",
                             tender_status="", stage="awards", object_id="a1",
                             bid_id="b1", lot_id="", event_time=None,
                             time_source="", time_basis="", time_reliability="HIGH",
                             object_status="unsuccessful", edrpou="12345678",
                             company_name="ТОВ «Х»", person=M.Person())]
        with self.assertRaises(M.MailError) as спроба:
            M.render_letter(події, "Добрий день, Іване!")
        self.assertEqual(спроба.exception.code, "EVENT_TIME_MISSING")

    def test_html_signature_links_to_the_site(self):
        html = self.лист()["body_html"]
        self.assertIn('href="https://tenderwin.in.ua"', html)
        self.assertIn('>TenderWin</a>', html)
        self.assertIn('href="tel:+380503101492"', html)

    def test_html_has_no_images_scripts_or_tracking(self):
        """Ні картинок, ні скриптів, ні пікселя відстеження."""
        html = self.лист()["body_html"]
        for заборонено in ("<img", "<script", "<iframe", "background-image",
                           "url(", "onclick", "<form"):
            self.assertNotIn(заборонено, html.lower(), заборонено)

    def test_dangerous_title_is_escaped(self):
        """Назва закупівлі з Prozorro не може зламати верстку листа."""
        row = self.лист(title='Ремонт <b>&</b> "покрівлі"')
        self.assertIn("&amp;", row["body_html"])
        self.assertNotIn("<b>", row["body_html"])
        self.assertIn("&lt;b&gt;", row["body_html"])

    def test_deadline_disappears_from_html_when_closed(self):
        row = self.лист(cp_end=iso(20, 0))          # строк уже минув
        self.assertNotIn("подати скаргу на це рішення можна до",
                         row["body_rendered"])
        self.assertNotIn("подати скаргу на це рішення можна до", row["body_html"])

    def test_text_only_mode_produces_no_html(self):
        M.LETTER_FORMAT = "TEXT_ONLY"
        row = self.лист()
        self.assertIsNone(row["body_html"])
        self.assertIn("Вартість — 3 499 грн", row["body_rendered"])

    def test_message_has_both_parts(self):
        """У листі дві частини: текст і HTML — і текст іде першим."""
        row = self.лист()
        raw = M.build_letter_bytes(row).decode("utf-8", "replace")
        self.assertIn("multipart/alternative", raw)
        self.assertLess(raw.index("text/plain"), raw.index("text/html"))

    def test_text_only_message_has_no_html_part(self):
        M.LETTER_FORMAT = "TEXT_ONLY"
        row = self.лист()
        raw = M.build_letter_bytes(row).decode("utf-8", "replace")
        self.assertNotIn("text/html", raw)

    def test_preview_writes_an_html_file_to_look_at(self):
        t = tender(bids=[bid("b1")],
                   awards=[award("a1", "b1", date=iso(22, 11), cp_start=iso(22, 11),
                                 cp_end=iso(27, 0))])
        eng, _ = self.run_scan([t])
        eng.build_letters()
        out = eng.preview()
        файли = os.listdir(out["тека"])
        eng.close()
        self.assertEqual(out["з_оформленням"], 1)
        self.assertTrue(any(f.endswith(".html") for f in файли))
        self.assertTrue(any(f.endswith(".eml") for f in файли))


class TestLetterV9(Base):
    """Лист V9 (06.10.2026): V8 з правками власника. Заморожена версія."""

    def setUp(self):
        super().setUp()
        M.LETTER_VERSION = "V9_OWNER"

    def лист(self, cp_end=None, title="Капітальний ремонт покрівлі"):
        t = tender(bids=[bid("b1")],
                   awards=[award("a1", "b1", date=iso(22, 11), cp_start=iso(22, 11),
                                 cp_end=cp_end or iso(27, 0))])
        t["title"] = title
        eng, _ = self.run_scan([t])
        eng.build_letters()
        row = eng.db.one("SELECT * FROM outreach")
        eng.close()
        return row

    def test_v9_still_renders_when_chosen(self):
        self.assertEqual(self.лист()["template_version"], "V9_OWNER")

    def test_text_matches_the_approved_wording(self):
        """Золотий текст V9 — слово в слово, як затвердив власник 06.10.2026."""
        row = self.лист()
        очікувано = (
            "Добрий день, Іване!\n"
            "\n"
            "Побачив у Prozorro, що пропозицію ТОВ «БУДІНВЕСТ» "
            "(ЄДРПОУ 12345678) в закупівлі UA-2026-09-01-000001-a "
            "«Капітальний ремонт покрівлі» відхилено "
            "(22 вересня 2026 року об 11:00).\n"
            "\n"
            "Питання після такого рішення зазвичай одне: чи справді підстави "
            "були і чи варто витрачати гроші на оскарження.\n"
            "\n"
            "Дам на нього дуже детальну письмову відповідь:\n"
            "\n"
            "• чи були підстави для відхилення,\n"
            "• які ризики,\n"
            "• що доцільно робити далі.\n"
            "\n"
            "Я зіставлю те, що написав Замовник у рішенні, з вимогами "
            "тендерної документації, документами Вашої тендерної пропозиції "
            "та практикою АМКУ. Якщо підстав для скарги немає, напишу про це "
            "прямо.\n"
            "\n"
            "Якщо підстави для скарги є, то складеного мною аналізу зазвичай "
            "Вам буде достатньо, щоб самостійно подати скаргу в АМКУ, не "
            "витрачаючи зайвих коштів на юридичний супровід.\n"
            "\n"
            "Вартість — 3 499 грн за одне рішення про відхилення, з усіма "
            "підставами, які в ньому названі.\n"
            "Висновок — протягом 24 годин після оплати.\n"
            "\n"
            "У Prozorro подати скаргу на це рішення можна до "
            "27 вересня 2026 року о 00:00.\n"
            "\n"
            "Побачити, як виглядає мій аналіз, Ви можете тут: "
            "https://tenderwin.in.ua\n"
            "\n"
            "З повагою,\n"
            "Віталій Щасливий\n"
            "радник з публічних закупівель\n"
            "+380 800 357 135\n")
        self.assertEqual(row["body_rendered"], очікувано)
        self.assertEqual(row["subject_rendered"],
                         "Відхилення у закупівлі UA-2026-09-01-000001-a: "
                         "чи є підстави для оскарження?")

    def test_html_has_the_same_words_as_the_text(self):
        import re as _re
        html = self.лист()["body_html"]
        текст_з_html = _re.sub(r"<[^>]+>", " ", html)
        текст_з_html = _re.sub(r"\s+", " ", текст_з_html).replace("&amp;", "&")
        for речення in ("що доцільно робити далі.",
                        "Я зіставлю те, що написав Замовник у рішенні, з вимогами "
                        "тендерної документації, документами Вашої тендерної "
                        "пропозиції та практикою АМКУ.",
                        "Якщо підстави для скарги є, то складеного мною аналізу "
                        "зазвичай Вам буде достатньо, щоб самостійно подати "
                        "скаргу в АМКУ, не витрачаючи зайвих коштів",
                        "Вартість — 3 499 грн",
                        "Побачити, як виглядає мій аналіз, Ви можете тут: "
                        "tenderwin.in.ua",
                        "З повагою, Віталій Щасливий радник з публічних "
                        "закупівель +380 800 357 135"):
            self.assertIn(речення, текст_з_html, речення)

    def test_old_wording_is_gone(self):
        row = self.лист()
        for було in ("Якщо питання актуальне", "відповідайте на цей лист",
                     "Я уважно зіставлю", "Зазвичай цього достатньо",
                     "(якщо це доцільно)", "вашими документами",
                     "50 310 14 92", "TenderWin"):
            self.assertNotIn(було, row["body_rendered"], було)
            self.assertNotIn(було, row["body_html"], було)

    def test_only_link_is_the_example_of_analysis(self):
        """Підпис без посилань: у листі одне посилання — на приклад аналізу."""
        html = self.лист()["body_html"]
        self.assertEqual(html.count("<a "), 1)
        self.assertIn('Ви можете тут: <a href="https://tenderwin.in.ua"', html)
        self.assertNotIn("tel:", html)
        підпис = html[html.index("З повагою"):]
        self.assertNotIn("<a", підпис)
        self.assertIn("радник з публічних закупівель<br>+380 800 357 135</p>",
                      підпис)

    def test_example_link_follows_the_setting(self):
        M.ANALYSIS_EXAMPLE_URL = "https://tenderwin.in.ua/pryklad"
        M.ANALYSIS_EXAMPLE_LABEL = "tenderwin.in.ua/pryklad"
        try:
            row = self.лист()
        finally:
            M.ANALYSIS_EXAMPLE_URL = M.SITE_URL
            M.ANALYSIS_EXAMPLE_LABEL = M.SITE_LABEL
        self.assertIn("тут: https://tenderwin.in.ua/pryklad\n", row["body_rendered"])
        self.assertIn('href="https://tenderwin.in.ua/pryklad"', row["body_html"])

    def test_two_rejections_work_in_v9(self):
        t1 = tender(uid="t1", ua="UA-2026-09-01-000001-a",
                    bids=[bid("b1")],
                    awards=[award("a1", "b1", date=iso(22, 11), cp_start=iso(22, 11),
                                  cp_end=iso(27, 0))])
        t2 = tender(uid="t2", ua="UA-2026-09-02-000002-a", value=9_000_000,
                    bids=[bid("b1")],
                    awards=[award("a2", "b1", date=iso(22, 15, 45),
                                  cp_start=iso(22, 15, 45), cp_end=iso(28, 0))])
        eng, _ = self.run_scan([t1, t2])
        eng.build_letters()
        row = eng.db.one("SELECT * FROM outreach")
        eng.close()
        self.assertEqual(row["template_version"], "V9_OWNER")
        self.assertIn("• UA-2026-09-02-000002-a", row["body_rendered"])
        self.assertIn("найраніший строк", row["body_rendered"])
        self.assertIn("+380 800 357 135", row["body_rendered"])

    def test_text_only_mode_produces_no_html(self):
        M.LETTER_FORMAT = "TEXT_ONLY"
        row = self.лист()
        self.assertIsNone(row["body_html"])
        self.assertIn("https://tenderwin.in.ua", row["body_rendered"])

    def test_html_has_no_images_scripts_or_tracking(self):
        html = self.лист()["body_html"]
        for заборонено in ("<img", "<script", "<iframe", "background-image",
                           "url(", "onclick", "<form"):
            self.assertNotIn(заборонено, html.lower(), заборонено)


class TestGmailAccess(Base):
    """Токен у файлі на Диску має знаходитись так само, як секрет Colab."""

    def setUp(self):
        super().setUp()
        self.збережений_TOKEN_DIR = E.TOKEN_DIR
        self.збережений_ALLOW = E.ALLOW_COLAB_SECRETS
        E.ALLOW_COLAB_SECRETS = False
        os.environ.pop(E.GMAIL_OAUTH_ENV, None)

    def tearDown(self):
        E.TOKEN_DIR = self.збережений_TOKEN_DIR
        E.ALLOW_COLAB_SECRETS = self.збережений_ALLOW
        os.environ.pop(E.GMAIL_OAUTH_ENV, None)
        super().tearDown()

    def test_token_file_next_to_the_database_is_found(self):
        диск = os.path.join(self.tmp, "Диск")
        os.makedirs(диск, exist_ok=True)
        with open(os.path.join(диск, E.TOKEN_FILE_NAME), "w", encoding="utf-8") as fh:
            fh.write('{"client_id":"x","client_secret":"y","refresh_token":"z"}')
        є, звідки = M.gmail_secret_ready(диск)
        self.assertTrue(є)
        self.assertIn(E.TOKEN_FILE_NAME, звідки)

    def test_send_without_argument_still_sees_the_token(self):
        """M.send() без dysk не має втрачати теку токена."""
        диск = os.path.join(self.tmp, "Диск2")
        os.makedirs(диск, exist_ok=True)
        with open(os.path.join(диск, E.TOKEN_FILE_NAME), "w", encoding="utf-8") as fh:
            fh.write('{"client_id":"x","client_secret":"y","refresh_token":"z"}')
        M.gmail_secret_ready(диск)      # так робить M.setup(DYSK)
        E.TOKEN_DIR = ""                 # новий сеанс Python, тека забулася
        M.DB_PATH = os.path.join(диск, "tenderwin_lead_mail.db")
        є, звідки = M.gmail_secret_ready()
        self.assertTrue(є, звідки)

    def test_missing_token_says_what_exactly_is_missing(self):
        порожня = os.path.join(self.tmp, "Порожня")
        os.makedirs(порожня, exist_ok=True)
        є, звідки = M.gmail_secret_ready(порожня)
        self.assertFalse(є)
        self.assertIn(E.GMAIL_OAUTH_ENV, звідки)
        self.assertIn(E.TOKEN_FILE_NAME, звідки)

    def test_environment_variable_wins(self):
        os.environ[E.GMAIL_OAUTH_ENV] = "{}"
        є, звідки = M.gmail_secret_ready(self.tmp)
        self.assertTrue(є)
        self.assertEqual(звідки, "змінна середовища")


class TestSchemaMigration(Base):
    def test_database_of_version_one_gets_the_new_column(self):
        """Стара база отримує колонку body_html і не втрачає рядків."""
        eng, _ = self.run_scan([tender(bids=[bid("b1")],
                                       awards=[award("a1", "b1", date=iso(22, 11),
                                                     cp_start=iso(22, 11))])])
        eng.build_letters()
        було = eng.db.one("SELECT COUNT(*) n FROM outreach")["n"]
        # відкочуємо базу до вигляду версії 1
        eng.db.x("UPDATE schema_meta SET value = '1' WHERE key = 'version'")
        eng.db.x("ALTER TABLE outreach DROP COLUMN body_html")
        eng.close()
        знову = M.LeadMail(db_path=self.db_path,
                           out_dir=os.path.join(self.tmp, "vyhid"), verbose=False)
        cols = {r["name"] for r in знову.db.q("PRAGMA table_info(outreach)")}
        стало = знову.db.one("SELECT COUNT(*) n FROM outreach")["n"]
        версія = знову.db.one(
            "SELECT value FROM schema_meta WHERE key = 'version'")["value"]
        знову.close()
        self.assertIn("body_html", cols)
        self.assertEqual(стало, було)
        self.assertEqual(версія, "2")


class TestDeduplication(Base):
    def legacy_db(self, edrpou="11111111", origin="ENGINE", status="SENT"):
        path = os.path.join(self.tmp, "leads.db")
        con = sqlite3.connect(path)
        con.executescript("""
            CREATE TABLE companies (company_id TEXT PRIMARY KEY,
                                    edrpou_norm TEXT, company_name TEXT);
            CREATE TABLE outreach_events (outreach_id TEXT PRIMARY KEY,
                company_id TEXT, outreach_type TEXT, mode TEXT, status TEXT,
                origin TEXT, generated_at TEXT, delivery_email_actual TEXT,
                contact_email_original TEXT);
            CREATE TABLE suppressions (suppression_id TEXT PRIMARY KEY,
                company_id TEXT, email TEXT, state TEXT, suppression_reason TEXT);
        """)
        con.execute("INSERT INTO companies VALUES ('c1', ?, 'ТОВ «БУДІНВЕСТ»')",
                    (edrpou,))
        con.execute("INSERT INTO outreach_events VALUES"
                    " ('o1','c1','FIRST_TOUCH','LIVE',?,?, '2026-09-01T10:00:00',"
                    "  'real@client.example','real@client.example')",
                    (status, origin))
        con.commit()
        con.close()
        return path

    def test_history_of_previous_engine_blocks_a_new_letter(self):
        M.LEGACY_DB_PATH = self.legacy_db()
        t = tender(bids=[bid("b1", code="11111111")],
                   awards=[award("a1", "b1", date=iso(22, 11), cp_start=iso(22, 11),
                                 code="11111111")])
        eng, _ = self.run_scan([t])
        stats = eng.build_letters()
        state = eng.db.one("SELECT state, state_reason FROM events")
        eng.close()
        self.assertEqual(stats["листів"], 0)
        self.assertEqual(state["state"], "NOT_ELIGIBLE")
        self.assertIn("вже отримувала", state["state_reason"])

    def test_rows_migrated_from_monolith_keep_their_own_status(self):
        M.LEGACY_DB_PATH = self.legacy_db(origin="MIGRATED")
        eng = self.engine()
        row = eng.db.one("SELECT status FROM outreach")
        eng.close()
        self.assertEqual(row["status"], "MIGRATED_UNVERIFIED",
                         "історія моноліту писалась ДО відправки — це не доказ")

    def test_suppression_blocks_the_company(self):
        eng = self.engine()
        t = tender(bids=[bid("b1", code="11111111")],
                   awards=[award("a1", "b1", date=iso(22, 11), cp_start=iso(22, 11),
                                 code="11111111")])
        cid = eng.repo.company("11111111", "ТОВ «БУДІНВЕСТ»")
        eng.repo.suppress(state="OPTED_OUT", reason="сказали «стоп»",
                          source="тест", company_id=cid)
        win = M.Window(*M.day_bounds(TODAY), mode="TODAY_ONLY")
        result = FakeApi([t]).scan(M.Window(win.start, TODAY, win.mode))
        eng.repo.start_run(eng.run_id, win)
        eng.ingest(result)
        stats = eng.build_letters()
        eng.close()
        self.assertEqual(stats["листів"], 0)
        self.assertEqual(stats["SUPPRESSED"], 1)

    def test_same_event_twice_is_not_duplicated(self):
        t = tender(bids=[bid("b1")], awards=[award("a1", "b1", date=iso(22, 11),
                                                   cp_start=iso(22, 11))])
        eng, _ = self.run_scan([t])
        eng.ingest(FakeApi([t]).scan(
            M.Window(*M.day_bounds(TODAY), mode="TODAY_ONLY")))
        count = eng.db.one("SELECT COUNT(*) n FROM events")["n"]
        eng.close()
        self.assertEqual(count, 1)


class TestTestMode(Base):
    def test_empty_allowlist_stops_everything(self):
        M.TEST_RECIPIENTS = []
        with self.assertRaises(M.MailError) as ctx:
            M.LeadMail(db_path=self.db_path, out_dir=self.tmp,
                       verbose=False).prepare()
        self.assertEqual(ctx.exception.code, "TEST_ALLOWLIST_EMPTY")

    def test_database_blocks_a_real_address(self):
        eng = self.engine()
        cid = eng.repo.company("11111111", "ТОВ")
        draft = M.Draft(run_id=eng.run_id, batch_id=eng.batch_id, company_id=cid,
                        contact_id=None, template_version="V4_OWNER",
                        subject="тема", body="тіло", html="",
                        contact_email="real@client.example",
                        delivery_email="real@client.example",
                        message_id="<x@tenderwin.com.ua>", event_ids=[])
        with self.assertRaises(M.MailError) as ctx:
            eng.repo.reserve(draft)
        eng.close()
        self.assertEqual(ctx.exception.code, "TEST_RECIPIENT_NOT_ALLOWED")

    def test_delivery_goes_only_to_allowlisted_boxes(self):
        eng, _ = self.run_scan([tender(bids=[bid("b1")],
                                       awards=[award("a1", "b1", date=iso(22, 11),
                                                     cp_start=iso(22, 11))])])
        eng.build_letters()
        row = eng.db.one("SELECT delivery_email_actual, contact_email_original"
                         " FROM outreach")
        eng.close()
        self.assertIn(row["delivery_email_actual"], M.TEST_RECIPIENTS)
        self.assertEqual(row["contact_email_original"], "office@bud.example")

    def test_address_cannot_be_swapped_by_an_update(self):
        """Правка рядка в базі не може підмінити адресу доставки на клієнтську."""
        eng, _ = self.run_scan([tender(bids=[bid("b1")],
                                       awards=[award("a1", "b1", date=iso(22, 11),
                                                     cp_start=iso(22, 11))])])
        eng.build_letters()
        with self.assertRaises(sqlite3.IntegrityError) as спроба:
            eng.db.x("UPDATE outreach SET delivery_email_actual = 'client@real.example'")
        eng.close()
        self.assertIn("DELIVERY_ADDRESS_IMMUTABLE", str(спроба.exception))

    def test_migrated_row_cannot_become_a_letter_to_a_client(self):
        """Рядок перенесеної історії не можна перевести в чергу з чужою адресою."""
        eng, _ = self.run_scan([tender(bids=[bid("b1")],
                                       awards=[award("a1", "b1", date=iso(22, 11),
                                                     cp_start=iso(22, 11))])])
        eng.db.x(
            "INSERT INTO outreach(outreach_id, company_id, contact_email_original,"
            " delivery_email_actual, mode, status, outreach_type, template_version,"
            " subject_rendered, body_rendered, message_id_header, event_ids,"
            " generated_at, batch_id, run_id) SELECT 'o-old', company_id,"
            " 'client@real.example', 'client@real.example', 'TEST',"
            " 'MIGRATED_UNVERIFIED', 'FIRST_TOUCH', 'LEGACY', '', '', 'mid-old',"
            " '[]', ?, 'B-old', 'R-old' FROM companies LIMIT 1",
            (M.now_iso(),))
        with self.assertRaises(sqlite3.IntegrityError) as спроба:
            eng.db.x("UPDATE outreach SET status = 'QUEUED' WHERE outreach_id = 'o-old'")
        eng.close()
        self.assertIn("TEST_RECIPIENT_NOT_ALLOWED", str(спроба.exception))

    def test_guards_are_added_to_a_database_of_an_older_build(self):
        """База без нових тригерів дістає їх під час відкриття, без втрати даних."""
        eng, _ = self.run_scan([tender(bids=[bid("b1")],
                                       awards=[award("a1", "b1", date=iso(22, 11),
                                                     cp_start=iso(22, 11))])])
        eng.build_letters()
        eng.db.x("DROP TRIGGER trg_delivery_address_immutable")
        eng.db.x("DROP TRIGGER trg_test_mode_fail_closed_update")
        було = eng.db.one("SELECT COUNT(*) n FROM outreach")["n"]
        eng.close()
        знову = M.LeadMail(db_path=self.db_path,
                           out_dir=os.path.join(self.tmp, "vyhid"), verbose=False)
        тригери = {r["name"] for r in знову.db.q(
            "SELECT name FROM sqlite_master WHERE type = 'trigger'")}
        стало = знову.db.one("SELECT COUNT(*) n FROM outreach")["n"]
        with self.assertRaises(sqlite3.IntegrityError):
            знову.db.x("UPDATE outreach SET delivery_email_actual = 'client@real.example'")
        знову.close()
        self.assertIn("trg_delivery_address_immutable", тригери)
        self.assertIn("trg_test_mode_fail_closed_update", тригери)
        self.assertEqual(стало, було)

    def test_delivery_address_is_stable_between_runs(self):
        first = M.delivery_address("company-1", M.TEST_RECIPIENTS)
        second = M.delivery_address("company-1", M.TEST_RECIPIENTS)
        self.assertEqual(first, second)

    def test_live_mode_still_needs_mail_access(self):
        """З 1.5.0 бойовий режим є, але без доступу до пошти транспорту немає."""
        M.MODE = "LIVE"
        with self.assertRaises(M.MailError) as ctx:
            M.make_transport(dry_run=False, out_dir=self.tmp)
        self.assertEqual(ctx.exception.code, "TRANSPORT_UNAVAILABLE")

    def test_unknown_mode_stops(self):
        M.MODE = "PROD"
        with self.assertRaises(M.MailError) as ctx:
            M.make_transport(dry_run=False, out_dir=self.tmp)
        self.assertEqual(ctx.exception.code, "MODE_UNKNOWN")


class TestSending(Base):
    def prepared(self):
        eng, _ = self.run_scan([tender(bids=[bid("b1")],
                                       awards=[award("a1", "b1", date=iso(22, 11),
                                                     cp_start=iso(22, 11))])])
        eng.build_letters()
        return eng

    def test_dry_run_writes_files_and_keeps_the_queue(self):
        eng = self.prepared()
        out = eng.preview()
        status = eng.db.one("SELECT status FROM outreach")["status"]
        eng.close()
        self.assertEqual(out["складено"], 1)
        self.assertEqual(status, "QUEUED")
        files = os.listdir(out["тека"])
        self.assertTrue(any(f.endswith(".eml") for f in files))
        self.assertTrue(any(f.endswith(".txt") for f in files))

    def test_successful_send_with_sent_label_is_confirmed(self):
        eng = self.prepared()
        stats = eng.send(transport=FakeTransport(), paced=False)
        row = eng.db.one("SELECT status, provider_message_id FROM outreach")
        eng.close()
        self.assertEqual(stats["SENT_CONFIRMED"], 1)
        self.assertEqual(row["status"], "SENT_CONFIRMED")
        self.assertEqual(row["provider_message_id"], "gmail-1")

    def test_without_sent_label_letter_is_only_accepted(self):
        eng = self.prepared()
        stats = eng.send(transport=FakeTransport(labels=("DRAFT",)), paced=False)
        row = eng.db.one("SELECT status, error_code FROM outreach")
        eng.close()
        self.assertEqual(stats["SENT"], 1)
        self.assertEqual(row["status"], "SENT")
        self.assertEqual(row["error_code"], "NOT_VERIFIED")

    def test_failure_keeps_the_lead_for_a_retry(self):
        """Відмова Gmail: лист лишається в черзі; після MAX_SEND_ATTEMPTS — SEND_FAILED."""
        eng = self.prepared()

        class Відмова(FakeTransport):
            def find_by_message_id(self, message_id_header):
                return None              # відхилений лист у «Надісланих» не лежить
        відмова = Відмова(FakeResult(ok=False, error="400"))
        стани = []
        for _ in range(M.MAX_SEND_ATTEMPTS):
            eng.send(transport=відмова, paced=False)
            стани.append(eng.db.one("SELECT status FROM outreach")["status"])
        eng.close()
        self.assertEqual(стани, ["QUEUED"] * (M.MAX_SEND_ATTEMPTS - 1) + ["SEND_FAILED"])

    def test_unknown_result_never_resends_blindly(self):
        eng = self.prepared()
        eng.send(transport=FakeTransport(FakeResult(ok=False, unknown=True,
                                                    error="timeout")), paced=False)
        status = eng.db.one("SELECT status FROM outreach")["status"]
        queued_again = eng.repo.queued(eng.batch_id, 10)
        eng.close()
        self.assertEqual(status, "DELIVERY_UNKNOWN")
        self.assertEqual(queued_again, [])

    def test_reconcile_finds_the_letter(self):
        eng = self.prepared()
        transport = FakeTransport(FakeResult(ok=False, unknown=True))
        row = eng.repo.queued(eng.batch_id, 1)[0]
        transport.sent.append((row["message_id_header"], b""))
        eng.send(transport=transport, paced=False)
        stats = eng.reconcile(transport=transport)
        status = eng.db.one("SELECT status FROM outreach")["status"]
        eng.close()
        self.assertEqual(stats["знайдено"], 1)
        self.assertEqual(status, "SENT_CONFIRMED")

    def test_reconcile_error_is_not_confused_with_not_found(self):
        eng = self.prepared()
        eng.send(transport=FakeTransport(FakeResult(ok=False, unknown=True)),
                 paced=False)
        stats = eng.reconcile(transport=FakeTransport(raise_on_find=True))
        row = eng.db.one("SELECT error_code FROM outreach")
        eng.close()
        self.assertEqual(stats["звірка_не_вдалась"], 1)
        self.assertEqual(row["error_code"], "RECONCILE_ERROR")

    def test_letters_of_an_older_batch_are_cancelled(self):
        eng = self.prepared()
        eng.db.x("UPDATE outreach SET batch_id = 'B_вчора'")
        eng.batch_id = "B_сьогодні"
        cancelled = eng.cancel_stale()
        row = eng.db.one("SELECT status, error_code FROM outreach")
        eng.close()
        self.assertEqual(cancelled, 1)
        self.assertEqual(row["status"], "CANCELLED")
        self.assertEqual(row["error_code"], "STALE_BATCH")


class TestStuckLetters(Base):
    """Обрив під час відправки: лист не губиться і не йде вдруге наосліп."""

    def стукнутий_лист(self):
        eng, _ = self.run_scan([tender(bids=[bid("b1")],
                                       awards=[award("a1", "b1", date=iso(22, 11),
                                                     cp_start=iso(22, 11))])])
        eng.build_letters()
        row = eng.db.one("SELECT * FROM outreach WHERE status = 'QUEUED'")
        eng.repo.set_outreach(row["outreach_id"], "SENDING")
        eng.close()
        return row

    def test_empty_queue_still_reconciles_a_stuck_letter(self):
        self.стукнутий_лист()
        транспорт = FakeTransport()
        збережено = M.make_transport
        M.make_transport = lambda dry_run, out_dir: транспорт
        try:
            підсумок = M.send()
        finally:
            M.make_transport = збережено
        двигун = M.LeadMail(db_path=self.db_path,
                            out_dir=os.path.join(self.tmp, "vyhid"), verbose=False)
        стан = двигун.db.one("SELECT status FROM outreach")["status"]
        двигун.close()
        self.assertEqual(підсумок["надіслано"], 0)
        self.assertEqual(підсумок["звірка"]["знайдено"], 0)
        self.assertEqual(стан, "DELIVERY_UNKNOWN",
                         "лист, якого немає в Gmail, лишається на ручний перегляд")
        self.assertEqual(транспорт.sent, [], "сліпого повтору бути не може")

    def test_reconcile_confirms_a_letter_that_did_leave(self):
        row = self.стукнутий_лист()
        транспорт = FakeTransport()
        транспорт.sent.append((row["message_id_header"], b""))   # лист таки пішов
        збережено = M.make_transport
        M.make_transport = lambda dry_run, out_dir: транспорт
        try:
            M.send()
        finally:
            M.make_transport = збережено
        двигун = M.LeadMail(db_path=self.db_path,
                            out_dir=os.path.join(self.tmp, "vyhid"), verbose=False)
        стан = двигун.db.one("SELECT status, provider_message_id p FROM outreach")
        двигун.close()
        self.assertEqual(стан["status"], "SENT_CONFIRMED")
        self.assertEqual(стан["p"], "gmail-1")


class TestScanCompleteness(Base):
    def test_feed_error_makes_the_run_incomplete(self):
        t = tender(bids=[bid("b1")], awards=[award("a1", "b1", date=iso(22, 11),
                                                   cp_start=iso(22, 11))])
        api = FakeApi([t], feed_pages_fail_at=2)
        api.store["t1"]["id"] = "t1"
        result = api.scan(M.Window(*M.day_bounds(TODAY), mode="TODAY_ONLY"))
        self.assertEqual(result.state, "INCOMPLETE")
        self.assertIn("стрічку прочитано не до кінця", result.note)

    def test_unreadable_tender_is_recorded_as_a_gap(self):
        t = tender(bids=[bid("b1")], awards=[award("a1", "b1", date=iso(22, 11),
                                                   cp_start=iso(22, 11))])
        eng, result = self.run_scan([t], broken={"t1"})
        gaps = eng.db.q("SELECT tender_uid, error_code FROM scan_gaps")
        eng.close()
        self.assertEqual(result.state, "INCOMPLETE")
        self.assertEqual(gaps[0]["error_code"], "HTTP_404")

    def test_first_page_failure_is_a_typed_error(self):
        api = FakeApi([], feed_pages_fail_at=1)
        with self.assertRaises(M.MailError) as ctx:
            api.scan(M.Window(*M.day_bounds(TODAY), mode="TODAY_ONLY"))
        self.assertEqual(ctx.exception.code, "PROZORRO_UNAVAILABLE")


class TestCardsAndNoDownloads(Base):
    def test_card_is_written_with_links_only(self):
        eng, _ = self.run_scan([tender(bids=[bid("b1")],
                                       awards=[award("a1", "b1", date=iso(22, 11),
                                                     cp_start=iso(22, 11))])])
        eng.build_letters()
        out = eng.cards()
        eng.close()
        self.assertEqual(out["карток"], 1)
        path = [os.path.join(dp, f) for dp, _, fs in os.walk(out["тека"])
                for f in fs if f.endswith(".docx")][0]
        import docx
        text = "\n".join(p.text for p in docx.Document(path).paragraphs)
        tables = "\n".join(c.text for t in docx.Document(path).tables
                           for r in t.rows for c in r.cells)
        self.assertIn("РОБОЧА КАРТКА ЗАКУПІВЛІ", text)
        self.assertIn("не завантажуються", text)
        self.assertIn("https://example/p1", tables)

    def test_existing_card_is_never_overwritten(self):
        eng, _ = self.run_scan([tender(bids=[bid("b1")],
                                       awards=[award("a1", "b1", date=iso(22, 11),
                                                     cp_start=iso(22, 11))])])
        eng.build_letters()
        first = eng.cards()
        path = [os.path.join(dp, f) for dp, _, fs in os.walk(first["тека"])
                for f in fs if f.endswith(".docx")][0]
        stamp = os.path.getmtime(path)
        with open(path, "ab") as fh:
            fh.write(b"")
        again = eng.cards()
        eng.close()
        self.assertEqual(again["карток"], 0, "нової картки не створено")
        self.assertEqual(again["уже_були"], 1, "наявна картка порахована окремо")
        self.assertEqual(os.path.getmtime(path), stamp)

    def test_module_never_downloads_documents(self):
        source = inspect.getsource(M)
        for forbidden in ("get_bytes(", "download_protocol(", "download_td(",
                          "urlretrieve", "iter_content"):
            self.assertNotIn(forbidden, source,
                             f"у скрипті не має бути завантаження: {forbidden}")

    def test_documents_are_kept_as_metadata(self):
        eng, result = self.run_scan([tender(bids=[bid("b1")],
                                            awards=[award("a1", "b1", date=iso(22, 11),
                                                          cp_start=iso(22, 11))])])
        row = eng.db.one("SELECT docs_decision, docs_bid, docs_td FROM events")
        eng.close()
        decision = json.loads(row["docs_decision"])
        self.assertEqual(decision[0]["назва"], "Протокол.pdf")
        self.assertTrue(decision[0]["посилання"].startswith("https://"))
        self.assertTrue(json.loads(row["docs_bid"]))
        self.assertTrue(json.loads(row["docs_td"]))


class TestCardsPerEvent(Base):
    def test_two_rejections_in_one_tender_get_two_cards(self):
        """Дві відмови тій самій компанії в одній закупівлі — дві картки."""
        t = tender(bids=[bid("b1"), bid("b2")],
                   awards=[award("a1", "b1", date=iso(22, 11), cp_start=iso(22, 11)),
                           award("a2", "b2", date=iso(22, 12), cp_start=iso(22, 12))])
        eng, _ = self.run_scan([t])
        eng.build_letters()
        stats = eng.cards()
        файли = []
        for корінь, _, назви in os.walk(stats["тека"]):
            файли += [n for n in назви if n.endswith(".docx")]
        eng.close()
        self.assertEqual(stats["карток"], 2)
        self.assertEqual(len(set(файли)), 2, "назви карток не можуть збігатися")

    def test_repeated_run_keeps_the_same_card_file(self):
        """Повторний прогін не створює другої картки і не чіпає нотаток."""
        t = tender(bids=[bid("b1")],
                   awards=[award("a1", "b1", date=iso(22, 11), cp_start=iso(22, 11))])
        eng, _ = self.run_scan([t])
        eng.build_letters()
        перший = eng.cards()
        шлях = [os.path.join(к, n) for к, _, нн in os.walk(перший["тека"])
                for n in нн if n.endswith(".docx")][0]
        with open(шлях, "ab") as fh:
            fh.write(b"")
        розмір = os.path.getsize(шлях)
        другий = eng.cards()
        eng.close()
        self.assertEqual(другий["карток"], 0, "картка вже є — другої не буде")
        self.assertEqual(os.path.getsize(шлях), розмір)


class TestReport(Base):
    def test_export_lists_every_event_with_its_state(self):
        t1 = tender(bids=[bid("b1", code="11111111")],
                    awards=[award("a1", "b1", date=iso(22, 11), cp_start=iso(22, 11),
                                  code="11111111")])
        t2 = tender(uid="t2", ua="UA-2026-09-02-000002-a",
                    bids=[bid("b2", code="22222222")],
                    awards=[award("a2", "b2", date=iso(22, 12), code="22222222")])
        eng, _ = self.run_scan([t1, t2])
        eng.build_letters()
        path = eng.export()
        eng.close()
        with open(path, encoding="utf-8-sig") as fh:
            body = fh.read()
        self.assertIn("UA-2026-09-01-000001-a", body)
        self.assertIn("NO_COMPLAINT_ROUTE", body)


if __name__ == "__main__":
    unittest.main(verbosity=2)
