# -*- coding: utf-8 -*-
"""
Приймальні тести TENDERWIN LEAD ENGINE — матриця §27.

Запуск:  python3 test_lead_engine.py
Мережа не потрібна: транспорт підмінюється, ліди приходять з фікстур.
"""
from __future__ import annotations

import json
import os
import shutil
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import tenderwin_lead_engine as E   # noqa: E402

# Тести не читають секрети середовища: інакше на машині з налаштованим
# GMAIL_OAUTH_JSON вони збудували б справжній клієнт Gmail і пішли в мережу.
E.ALLOW_COLAB_SECRETS = False


# ---------------------------------------------------------------- допоміжне
class FailingTransport:
    """Gmail відмовляє. Лист точно не пішов."""
    name = "FAILING"

    def send(self, raw, mid):
        return E.SendResult(ok=False, error="500 backend error")

    def find_by_message_id(self, mid):
        return None


class HangingTransport:
    """Обрив звʼязку. Лист МІГ піти — це не те саме, що відмова."""
    name = "HANGING"

    def __init__(self, findable=False):
        self.findable = findable

    def send(self, raw, mid):
        return E.SendResult(ok=False, unknown=True, error="connection reset")

    def find_by_message_id(self, mid):
        return {"id": "gmail-123", "threadId": "thr-1"} if self.findable else None


class FakeGmailTransport(E.DryRunTransport):
    """
    Поводиться як налаштований провайдер (щоб preflight не блокував за
    транспортом), але нічого не надсилає. Дозволяє тестувати бойовий шлях.
    """
    name = "GMAIL"


class CapturingTransport(E.DryRunTransport):
    """Запамʼятовує сирі листи, щоб перевірити заголовки."""

    def __init__(self, out_dir):
        super().__init__(out_dir)
        self.raw = []

    def send(self, raw_rfc822, message_id_header):
        self.raw.append(raw_rfc822)
        return super().send(raw_rfc822, message_id_header)


def facts(edrpou="12345678", ua="UA-2026-08-28-000001-a", email="klient@example.com",
          name="Іван Петрович Коваленко", company="ТОВ «Будінвест»",
          reason="Учасником не надано довідку про наявність працівників.",
          bid="bid-1", award="award-1"):
    return E.RejectionFacts(
        ua_id=ua, edrpou=edrpou, company_name=company,
        rejection_date="2026-08-28T14:00:00+03:00", reason_raw=reason,
        tender=E.TenderRef(tender_id="t-" + ua, ua_id=ua, title="Капітальний ремонт",
                           buyer="Замовник", cpv="45000000-7"),
        person=E.Person(name_raw=name, email=email, phone="+380501112233",
                        name_source="PROZORRO_ESZ"),
        bid_id=bid, award_id=award,
        source_locator={"source": "prozorro.award"})


class Base(unittest.TestCase):
    mode = "TEST"

    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.db_path = os.path.join(self.tmp, "leads.db")

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def engine(self, items, *, mode=None, transport=None, legacy=None,
               reason_mode="verbatim", db_path=None, gate_open=True,
               migrated=True):
        mode = mode or self.mode
        cfg = E.Config(db_path=db_path or self.db_path,
                       out_dir=os.path.join(self.tmp, "out"),
                       mode=mode, reason_mode=reason_mode,
                       live_send_enabled=(gate_open and mode == "LIVE"),
                       min_delay=0, max_delay=0, verbose=False)
        if transport is None:
            transport = (FakeGmailTransport(os.path.join(self.tmp, "eml"))
                         if mode == "LIVE"
                         else E.DryRunTransport(os.path.join(self.tmp, "eml")))
        eng = E.LeadEngine(cfg, E.FixtureSource(items), transport, legacy)
        if migrated:
            eng.repo.mark_history_migrated("тестова фікстура")
        return eng

    def run_cycle(self, eng, items):
        eng.ingest(items)
        return eng.build_queue()


# ================================================================ §27 матриця
class TestDeduplication(Base):
    """1, 2, 3, 4, 5, 18 — дедуплікація за компанією."""

    def test_01_same_edrpou_same_email_next_day(self):
        eng = self.engine([facts()], mode="LIVE")
        self.run_cycle(eng, [facts()])
        eng.send()
        eng.close()
        # наступного дня — та сама компанія, та сама адреса, інша закупівля
        eng2 = self.engine([], mode="LIVE")
        q = self.run_cycle(eng2, [facts(ua="UA-2026-08-29-000002-a", bid="bid-9",
                                        award="award-9")])
        eng2.close()
        self.assertEqual(q["QUEUED"], 0, "другий перший лист не має ставитись у чергу")

    def test_02_same_edrpou_different_email(self):
        eng = self.engine([], mode="LIVE")
        self.run_cycle(eng, [facts(email="a@x.ua")])
        eng.send()
        q = self.run_cycle(eng, [facts(email="b@x.ua", ua="UA-2026-08-29-000003-a",
                                       bid="b2", award="a2")])
        eng.close()
        self.assertEqual(q["QUEUED"], 0, "інша адреса тієї ж компанії — не новий лід")

    def test_03_same_edrpou_different_tender(self):
        eng = self.engine([], mode="LIVE")
        self.run_cycle(eng, [facts()])
        eng.send()
        q = self.run_cycle(eng, [facts(ua="UA-2026-08-30-000004-a", bid="b3",
                                       award="a3")])
        eng.close()
        self.assertEqual(q["QUEUED"], 0, "інша закупівля — теж не другий перший лист")

    def test_04_different_edrpou_same_test_inbox(self):
        eng = self.engine([])
        q = self.run_cycle(eng, [facts(edrpou="11111111", bid="b1", award="a1"),
                                 facts(edrpou="22222222",
                                       ua="UA-2026-08-28-000002-a",
                                       bid="b2", award="a2")])
        eng.close()
        self.assertEqual(q["QUEUED"], 2, "різні компанії на одну скриньку — обидві")

    def test_05_restart_after_sent(self):
        eng = self.engine([], mode="LIVE")
        self.run_cycle(eng, [facts()])
        eng.send()
        eng.close()
        eng2 = self.engine([], mode="LIVE")          # «перезапуск»
        q = self.run_cycle(eng2, [facts(ua="UA-2026-09-01-000009-a", bid="bx",
                                        award="ax")])
        eng2.close()
        self.assertEqual(q["QUEUED"], 0)

    def test_18_database_reopen_next_day(self):
        eng = self.engine([], mode="LIVE")
        self.run_cycle(eng, [facts()])
        eng.send()
        eng.close()
        eng2 = self.engine([], mode="LIVE", db_path=self.db_path)
        self.assertTrue(eng2.repo.has_first_touch(
            eng2.repo.get_or_create_company("12345678", "ТОВ «Будінвест»")))
        eng2.close()


class TestTestModeDoesNotBurn(Base):
    """Пастка 1: тестові листи не витрачають довічний перший дотик."""

    def test_test_send_does_not_consume_first_touch(self):
        eng = self.engine([], mode="TEST")
        q1 = self.run_cycle(eng, [facts()])
        eng.send()
        eng.close()
        self.assertEqual(q1["QUEUED"], 1)

        eng2 = self.engine([], mode="LIVE")
        q2 = self.run_cycle(eng2, [facts(ua="UA-2026-09-02-000010-a", bid="bz",
                                         award="az")])
        eng2.close()
        self.assertEqual(q2["QUEUED"], 1,
                         "після ТЕСТОВОГО листа бойовий перший лист має бути можливим")


class TestIdempotency(Base):
    """6, 15, 16 — повтор ніколи не створює другий лист."""

    def test_06_crash_during_sending_no_blind_resend(self):
        eng = self.engine([], mode="LIVE", transport=HangingTransport(findable=False))
        self.run_cycle(eng, [facts()])
        stats = eng.send()
        self.assertEqual(stats["DELIVERY_UNKNOWN"], 1)
        q = self.run_cycle(eng, [facts(ua="UA-2026-09-03-000011-a", bid="bq",
                                       award="aq")])
        eng.close()
        self.assertEqual(q["QUEUED"], 0, "DELIVERY_UNKNOWN блокує другий лист")

    def test_06b_reconcile_finds_message(self):
        eng = self.engine([], mode="LIVE", transport=HangingTransport(findable=True))
        self.run_cycle(eng, [facts()])
        eng.send()
        rec = eng.reconcile()
        eng.close()
        self.assertEqual(rec["SENT"], 1, "звірка знайшла лист -> статус SENT")

    def test_06c_reconcile_inconclusive_goes_manual(self):
        eng = self.engine([], mode="LIVE", transport=HangingTransport(findable=False))
        self.run_cycle(eng, [facts()])
        eng.send()
        rec = eng.reconcile()
        row = eng.db.one("SELECT status, error_code FROM outreach_events LIMIT 1")
        eng.close()
        self.assertEqual(rec["MANUAL"], 1)
        self.assertEqual(row["error_code"], "MANUAL_REVIEW_REQUIRED")

    def test_15_failed_send_allows_retry_without_duplicate(self):
        eng = self.engine([], mode="LIVE", transport=FailingTransport())
        self.run_cycle(eng, [facts()])
        stats = eng.send()
        self.assertEqual(stats["SEND_FAILED"], 1)
        n = eng.db.one("SELECT COUNT(*) n FROM outreach_events")["n"]
        eng.close()
        self.assertEqual(n, 1, "невдала відправка не створює другий рядок звернення")

    def test_15b_same_idempotency_key_twice(self):
        eng = self.engine([], mode="LIVE")
        self.run_cycle(eng, [facts()])
        row = eng.repo.queued("LIVE", 1)[0]
        key = row["message_id_header"].strip("<>")
        eng.repo.open_attempt(row["outreach_id"], key, 0)
        with self.assertRaises(E.EngineError) as ctx:
            eng.repo.open_attempt(row["outreach_id"], key, 1)
        eng.close()
        self.assertEqual(ctx.exception.code, "DUPLICATE_SEND_ATTEMPT")

    def test_16_successful_send_persists_provider_ids(self):
        eng = self.engine([], mode="LIVE")
        self.run_cycle(eng, [facts()])
        eng.send()
        row = eng.db.one("SELECT status, gmail_message_id, sent_at"
                         " FROM outreach_events LIMIT 1")
        eng.close()
        self.assertEqual(row["status"], "SENT")
        self.assertTrue(row["gmail_message_id"])
        self.assertTrue(row["sent_at"])


class TestSuppression(Base):
    """7 — стоп-лист сильніший за все."""

    def test_07_opted_out_company_gets_nothing(self):
        eng = self.engine([], mode="LIVE")
        eng.ingest([facts()])
        cid = eng.repo.get_or_create_company("12345678", "ТОВ «Будінвест»")
        eng.repo.suppress("OPTED_OUT", "відповів «стоп»", "gmail_reply",
                          company_id=cid)
        q = eng.build_queue()
        eng.close()
        self.assertEqual(q["QUEUED"], 0)
        self.assertEqual(q.get("SUPPRESSED"), 1)

    def test_07b_reply_stop_creates_suppression(self):
        eng = self.engine([], mode="LIVE")
        self.run_cycle(eng, [facts()])
        eng.send()
        oid = eng.db.one("SELECT outreach_id FROM outreach_events")["outreach_id"]
        kind = eng.replies.apply(oid, "Re: Щодо закупівлі", "Стоп, не пишіть більше.")
        n = eng.db.one("SELECT COUNT(*) n FROM suppressions")["n"]
        eng.close()
        self.assertEqual(kind, "OPTED_OUT")
        self.assertEqual(n, 1, "P.S. обіцяє «стоп» — обіцянка має виконуватись")

    def test_07c_meaningful_reply_recognised(self):
        self.assertEqual(
            E.classify_reply("Re: закупівля", "Так, подивіться будь ласка мою ситуацію"),
            "MEANINGFUL_REPLY")

    def test_07d_unclear_reply_goes_manual(self):
        self.assertEqual(E.classify_reply("Re:", "Доброго дня."),
                         "MANUAL_REVIEW_REQUIRED")

    def test_20_permanent_bounce_suppresses_email_not_company(self):
        eng = self.engine([], mode="LIVE")
        self.run_cycle(eng, [facts()])
        eng.send()
        oid = eng.db.one("SELECT outreach_id FROM outreach_events")["outreach_id"]
        eng.replies.apply_bounce(oid, permanent=True)
        row = eng.db.one("SELECT state, company_id, email FROM suppressions")
        eng.close()
        self.assertEqual(row["state"], "INVALID_EMAIL")
        self.assertIsNone(row["company_id"], "гасимо адресу, а не всю компанію")


class TestPersonalization(Base):
    """8, 9 — вигаданої персоналізації не буває."""

    def test_09_unknown_name_uses_safe_greeting(self):
        eng = self.engine([], mode="LIVE")
        self.run_cycle(eng, [facts(name="")])
        row = eng.db.one("SELECT body_rendered, contact_vocative FROM outreach_events"
                         " o JOIN contacts c ON c.contact_id=o.contact_id")
        eng.close()
        self.assertIn(E.SAFE_GREETING, row["body_rendered"])

    def test_08_no_reason_goes_manual_review(self):
        eng = self.engine([], mode="LIVE")
        q = self.run_cycle(eng, [facts(reason="")])
        eng.close()
        self.assertEqual(q["QUEUED"], 0)
        self.assertEqual(q.get("MANUAL_REVIEW_REQUIRED"), 1)

    def test_08b_no_email_goes_manual_review(self):
        eng = self.engine([], mode="LIVE")
        q = self.run_cycle(eng, [facts(email=None)])
        eng.close()
        self.assertEqual(q.get("MANUAL_REVIEW_REQUIRED"), 1)

    def test_08c_no_edrpou_is_not_ingested(self):
        eng = self.engine([], mode="LIVE")
        st = eng.ingest([facts(edrpou="")])
        eng.close()
        self.assertEqual(st["без_ЄДРПОУ"], 1, "e-mail і назва ключем бути не можуть")

    def test_classifier_reason_blocked_in_live(self):
        eng = self.engine([], mode="LIVE", reason_mode="classifier")
        q = self.run_cycle(eng, [facts()])
        eng.close()
        self.assertEqual(q["QUEUED"], 0,
                         "категорія класифікатора не допускається в бойовий лист")

    def test_verbatim_reason_is_quoted_exactly(self):
        text = "Учасником не надано довідку про відсутність заборгованості зі сплати податків."
        eng = self.engine([], mode="LIVE")
        self.run_cycle(eng, [facts(reason=text)])
        body = eng.db.one("SELECT body_rendered FROM outreach_events")["body_rendered"]
        eng.close()
        self.assertIn(text.rstrip("."), body, "у листі мають бути слова замовника")
        self.assertNotIn("пункт", body.lower(),
                         "жодних вигаданих посилань на норми в листі")


class TestExperiment(Base):
    """10 — призначення варіанта переживає перезапуск."""

    def test_10_assignment_is_deterministic(self):
        a = E.ExperimentAssigner.compute("EXP", "company-abc")
        b = E.ExperimentAssigner.compute("EXP", "company-abc")
        self.assertEqual(a, b)
        self.assertIn(a, ("A", "B"))

    def test_10b_assignment_survives_restart(self):
        eng = self.engine([], mode="LIVE")
        eng.ingest([facts()])
        cid = eng.repo.get_or_create_company("12345678", "ТОВ «Будінвест»")
        v1 = eng.experiment.variant_for(cid)
        eng.close()
        eng2 = self.engine([], mode="LIVE")
        v2 = eng2.experiment.variant_for(cid)
        eng2.close()
        self.assertEqual(v1, v2)

    def test_10c_allocation_is_balanced(self):
        counts = {"A": 0, "B": 0}
        for i in range(1000):
            counts[E.ExperimentAssigner.compute("EXP", f"c-{i}")] += 1
        self.assertTrue(400 < counts["A"] < 600, counts)

    def test_11_template_version_immutable(self):
        eng = self.engine([], mode="LIVE")
        bad = E.Template("TENDERWIN_FIRST_TOUCH_B", "V1", "B", "тема", "ІНШЕ ТІЛО")
        with self.assertRaises(E.EngineError) as ctx:
            eng.repo.upsert_template(bad)
        eng.close()
        self.assertEqual(ctx.exception.code, "TEMPLATE_VERSION_CONFLICT")


class TestTestModeSafety(Base):
    """12, 13 — тестовий режим падає закрито."""

    def test_12_real_client_email_hard_blocked_in_test(self):
        eng = self.engine([], mode="TEST")
        eng.ingest([facts()])
        cid = eng.repo.get_or_create_company("12345678", "ТОВ «Будінвест»")
        draft = E.OutreachDraft(
            company_id=cid, contact_id=None, lead_id=None,
            outreach_type="FIRST_TOUCH", mode="TEST",
            experiment_id="E", variant="B",
            template_id="TENDERWIN_FIRST_TOUCH_B", template_version="V1",
            subject="тема", body="тіло",
            contact_email_original="klient@example.com",
            delivery_email_actual="klient@example.com",   # реальна адреса!
            message_id_header=E.build_message_id("k1"))
        with self.assertRaises(E.EngineError) as ctx:
            eng.repo.reserve_outreach(draft)
        eng.close()
        self.assertEqual(ctx.exception.code, "TEST_RECIPIENT_NOT_ALLOWED")

    def test_13_no_cc_bcc_fields_exist_at_all(self):
        raw = E.build_rfc822(sender_name="X", sender_email="a@b.ua",
                             to_email="t@b.ua", subject="s", body="b",
                             message_id="<k@b.ua>", intended_to="real@c.ua")
        text = raw.decode("utf-8", "replace")
        self.assertNotIn("\nCc:", text)
        self.assertNotIn("\nBcc:", text)
        self.assertIn("X-TenderWin-Intended-To: real@c.ua", text)

    def test_13b_delivery_and_intended_stored_separately(self):
        eng = self.engine([], mode="TEST")
        self.run_cycle(eng, [facts()])
        row = eng.db.one("SELECT contact_email_original, delivery_email_actual"
                         " FROM outreach_events")
        eng.close()
        self.assertEqual(row["contact_email_original"], "klient@example.com")
        self.assertIn(row["delivery_email_actual"], E.TEST_RECIPIENTS)

    def test_11b_rotation_is_reproducible(self):
        r1 = E.TestRecipientRotation(E.TEST_RECIPIENTS)
        r2 = E.TestRecipientRotation(E.TEST_RECIPIENTS)
        self.assertEqual([r1.next() for _ in range(5)],
                         [r2.next() for _ in range(5)])


class TestLiveGate(Base):
    """14, 17 — бойовий режим і звіт."""

    def test_14_live_blocked_without_explicit_gate(self):
        eng = self.engine([], mode="LIVE", gate_open=False)
        rep = eng.preflight.check()
        eng.close()
        self.assertFalse(rep.passed)
        self.assertTrue(any("TEST_MODE" in b for b in rep.blockers))

    def test_14b_live_blocked_lists_unverified_items(self):
        eng = self.engine([], mode="LIVE", gate_open=False)
        rep = eng.preflight.check()
        eng.close()
        joined = " ".join(rep.unverified)
        self.assertIn("DKIM", joined)
        self.assertIn("Message-ID", joined)

    def test_14c_unmigrated_history_blocks_live(self):
        eng = self.engine([], mode="LIVE", migrated=False)
        rep = eng.preflight.check()
        eng.close()
        self.assertFalse(rep.passed)
        self.assertTrue(any("історія попередніх розсилок" in b for b in rep.blockers),
                        "без перенесеної історії — повторний «перший» лист усім")

    def test_14d_dry_run_transport_blocks_live(self):
        eng = self.engine([], mode="LIVE",
                          transport=E.DryRunTransport(os.path.join(self.tmp, "e2")))
        rep = eng.preflight.check()
        eng.close()
        self.assertTrue(any("DRY_RUN" in b for b in rep.blockers))

    def test_14e_gate_open_and_migrated_passes(self):
        eng = self.engine([], mode="LIVE")
        rep = eng.preflight.check()
        eng.close()
        self.assertTrue(rep.passed, rep.render())

    def test_17_export_failure_cannot_cause_resend(self):
        eng = self.engine([], mode="LIVE")
        self.run_cycle(eng, [facts()])
        eng.send()
        with self.assertRaises(Exception):
            eng.db.conn.execute("INSERT INTO v_outreach_export(mode) VALUES ('LIVE')")
        q = self.run_cycle(eng, [facts(ua="UA-2026-09-05-000012-a", bid="bb",
                                       award="aa")])
        eng.close()
        self.assertEqual(q["QUEUED"], 0)

    def test_17b_export_produces_full_row(self):
        eng = self.engine([], mode="TEST")
        self.run_cycle(eng, [facts()])
        eng.send()
        out = eng.export("test")
        eng.close()
        self.assertTrue(os.path.exists(out["csv"]))
        with open(out["csv"], encoding="utf-8-sig") as fh:
            head = fh.readline().strip().split(",")
        for col in ("ua_id", "delivery_email_actual", "mode", "variant",
                    "body_rendered", "gmail_message_id"):
            self.assertIn(col, head)

    def test_funnel_counts_live_only(self):
        eng = self.engine([], mode="TEST")
        self.run_cycle(eng, [facts()])
        eng.send()
        f = eng.exporter.funnel()
        eng.close()
        self.assertEqual(f["letters_sent_live"], 0)
        self.assertEqual(f["letters_sent_test"], 1)


class TestStorageGate(Base):
    """Немає бази — немає відправки."""

    def test_storage_unavailable_is_typed_and_stops(self):
        with self.assertRaises(E.EngineError) as ctx:
            E.Db("/proc/1/nonexistent/leads.db")
        self.assertEqual(ctx.exception.code, "STORAGE_UNAVAILABLE")
        self.assertFalse(ctx.exception.may_continue)

    def test_edrpou_normalisation(self):
        self.assertEqual(E.normalize_edrpou("12345678"), "12345678")
        self.assertEqual(E.normalize_edrpou("ЄДРПОУ 12345678"), "12345678")
        self.assertEqual(E.normalize_edrpou("1234567890"), "1234567890")
        self.assertEqual(E.normalize_edrpou("123456789"), "0123456789")
        self.assertIsNone(E.normalize_edrpou(""))
        self.assertIsNone(E.normalize_edrpou("нема"))


class TestRetryPath(Base):
    """Тимчасова помилка не має стирати лід назавжди."""

    def test_failed_send_is_retried_next_run(self):
        eng = self.engine([], mode="LIVE", transport=FailingTransport())
        self.run_cycle(eng, [facts()])
        first = eng.send()
        self.assertEqual(first["SEND_FAILED"], 1)
        eng.close()

        eng2 = self.engine([], mode="LIVE")          # наступного дня Gmail живий
        second = eng2.send()
        n = eng2.db.one("SELECT COUNT(*) n FROM outreach_events")["n"]
        eng2.close()
        self.assertEqual(second["повторно_в_черзі"], 1)
        self.assertEqual(second["SENT"], 1, "лід має дійти з другої спроби")
        self.assertEqual(n, 1, "і при цьому лишитись ОДНИМ листом")

    def test_retry_limit_cancels_instead_of_looping(self):
        eng = self.engine([], mode="LIVE", transport=FailingTransport())
        self.run_cycle(eng, [facts()])
        for _ in range(4):
            eng.send(max_retries=2)
        row = eng.db.one("SELECT status, error_code FROM outreach_events")
        eng.close()
        self.assertEqual(row["status"], "CANCELLED")
        self.assertEqual(row["error_code"], "RETRY_LIMIT_REACHED")

    def test_delivery_unknown_is_never_requeued(self):
        eng = self.engine([], mode="LIVE", transport=HangingTransport(findable=False))
        self.run_cycle(eng, [facts()])
        eng.send()
        again = eng.send()
        row = eng.db.one("SELECT status FROM outreach_events")
        eng.close()
        self.assertEqual(again["повторно_в_черзі"], 0,
                         "невідомий стан доставки не можна повторювати наосліп")
        self.assertEqual(row["status"], "DELIVERY_UNKNOWN")

    def test_cancelled_frees_company_for_a_future_letter(self):
        eng = self.engine([], mode="LIVE", transport=FailingTransport())
        self.run_cycle(eng, [facts()])
        for _ in range(3):
            eng.send(max_retries=1)
        cid = eng.repo.get_or_create_company("12345678", "ТОВ «Будінвест»")
        free = not eng.repo.has_first_touch(cid)
        eng.close()
        self.assertTrue(free, "лист, який так і не пішов, не має спалювати дотик")


class TestMigrationAndEdges(Base):
    """Крок 0 і кути, на яких зазвичай ламається."""

    def _csv(self, name, header, rows):
        path = os.path.join(self.tmp, name)
        with open(path, "w", encoding="utf-8-sig", newline="") as fh:
            fh.write(",".join(header) + "\n")
            for r in rows:
                fh.write(",".join(r) + "\n")
        return path

    def test_migration_marks_previous_recipients(self):
        hist = self._csv("istoriya.csv",
                         ["дата", "ключ", "ЄДРПОУ_або_ІПН", "компанія", "email"],
                         [["2026-07-01", "ed:12345678", "12345678", "ТОВ Тест",
                           "old@x.ua"]])
        stop = self._csv("stop.csv", ["ЄДРПОУ_або_ІПН", "компанія", "email", "причина"],
                         [["87654321", "ТОВ Стоп", "stop@x.ua", "просив не писати"]])
        st = E.migrate_legacy_history(self.db_path, history_csv=hist,
                                      stop_list_csv=stop)
        self.assertEqual(st["звернень"], 1)
        self.assertEqual(st["стопів"], 1)

        eng = self.engine([], mode="LIVE", migrated=False)
        q = self.run_cycle(eng, [facts(edrpou="12345678", ua="UA-2026-09-10-1-a",
                                       bid="bm", award="am"),
                                 facts(edrpou="87654321", ua="UA-2026-09-10-2-a",
                                       bid="bn", award="an")])
        self.assertTrue(eng.repo.is_history_migrated())
        eng.close()
        self.assertEqual(q["QUEUED"], 0,
                         "той, кому вже писали, і той, хто сказав «стоп», листа не отримують")

    def test_two_rejections_same_company_one_run_give_one_letter(self):
        eng = self.engine([], mode="LIVE")
        q = self.run_cycle(eng, [facts(ua="UA-2026-09-11-1-a", bid="b1", award="a1"),
                                 facts(ua="UA-2026-09-11-2-a", bid="b2", award="a2")])
        n = eng.db.one("SELECT COUNT(*) n FROM outreach_events")["n"]
        eng.close()
        self.assertEqual(q["QUEUED"], 1, "дві підстави в один день — один лист")
        self.assertEqual(n, 1)

    def test_company_name_with_braces_does_not_break_render(self):
        eng = self.engine([], mode="LIVE")
        q = self.run_cycle(eng, [facts(company="ТОВ {Альфа}")])
        body = eng.db.one("SELECT body_rendered FROM outreach_events")["body_rendered"]
        eng.close()
        self.assertEqual(q["QUEUED"], 1)
        self.assertIn("ТОВ {Альфа}", body)

    def test_same_rejection_ingested_twice_is_not_duplicated(self):
        eng = self.engine([], mode="LIVE")
        eng.ingest([facts()])
        st = eng.ingest([facts()])
        n = eng.db.one("SELECT COUNT(*) n FROM rejections")["n"]
        eng.close()
        self.assertEqual(st["вже_були"], 1)
        self.assertEqual(n, 1, "та сама подія відхилення не подвоюється")

    def test_migrated_rows_do_not_pollute_metrics(self):
        """Перенесені історичні листи надсилали не ми — у метриках їм не місце."""
        hist = self._csv("h.csv", ["дата", "ЄДРПОУ_або_ІПН", "компанія", "email"],
                         [["2026-07-01", "55555555", "ТОВ Старий", "o@x.ua"]])
        E.migrate_legacy_history(self.db_path, history_csv=hist)
        eng = self.engine([], mode="TEST", migrated=False)
        self.run_cycle(eng, [facts()])
        eng.send()
        f = eng.exporter.funnel()
        eng.close()
        self.assertEqual(f["letters_sent_live"], 0,
                         "бойових листів движок ще не надсилав")
        self.assertEqual(f["letters_sent_test"], 1)
        self.assertEqual(f["перенесено_з_історії"], 1)
        self.assertIsNone(f["meaningful_reply_rate"],
                          "частка відгуку не рахується від чужих листів")

    @staticmethod
    def _schema_v1_sql() -> str:
        """Схема такою, якою вона була у версії 1: без origin, без полів картки,
        зі старим індексом ux_first_touch_live (предикат mode='LIVE')."""
        # колонки, доданих у v2 та v4, у схемі v1 не було
        added = ("origin", "stage", "lot_amount", "complaint_deadline",
                 "contact_region", "contact_address")
        out, skip_next = [], False
        for line in E.SCHEMA_SQL.splitlines(keepends=True):
            head = line.strip().split(" ")[0]
            if skip_next and line.lstrip().startswith("CHECK ("):
                continue                      # продовження визначення origin
            skip_next = head in added
            if skip_next:
                continue
            out.append(line)
        sql = "".join(out)
        import re as _re
        for col in added:
            assert not _re.search(rf"\b{col}\b", sql), \
                f"схема змінилася: {col} лишився у v1-симуляції"
        i = sql.index("CREATE UNIQUE INDEX ux_first_touch")
        j = sql.index(";", i) + 1
        return sql[:i] + (
            "CREATE UNIQUE INDEX ux_first_touch_live"
            " ON outreach_events(company_id, outreach_type)"
            " WHERE mode = 'LIVE' AND status IN ('QUEUED','SENDING','SENT',"
            "'DELIVERY_UNKNOWN','BOUNCED','REPLIED','OPTED_OUT');") + sql[j:]

    def test_schema_v1_database_is_migrated_not_rejected(self):
        """Стара база не має падати з помилкою — вона має мігрувати до кінця."""
        import sqlite3
        con = sqlite3.connect(self.db_path)
        con.executescript(self._schema_v1_sql())
        con.execute("INSERT INTO schema_meta(key,value) VALUES ('version','1')")
        con.commit()
        con.close()
        db = E.Db(self.db_path)          # має мігрувати, а не кинути виняток
        v = db.one("SELECT value FROM schema_meta WHERE key='version'")["value"]
        cols = [r[1] for r in db.conn.execute("PRAGMA table_info(outreach_events)")]
        rej = [r[1] for r in db.conn.execute("PRAGMA table_info(rejections)")]
        ct = [r[1] for r in db.conn.execute("PRAGMA table_info(contacts)")]
        idx = {r[0]: r[1] for r in db.conn.execute(
            "SELECT name, sql FROM sqlite_master WHERE type='index'"
            " AND name LIKE 'ux_first_touch%'")}
        db.close()
        self.assertEqual(v, str(E.SCHEMA_VERSION), "міграція не дійшла до кінця")
        self.assertIn("origin", cols)                       # v2
        self.assertIn("ux_first_touch", idx)                # v3
        self.assertNotIn("ux_first_touch_live", idx, "старий індекс лишився")
        self.assertIn("mode", idx["ux_first_touch"],
                      "новий індекс має включати mode у ключ")
        for c in ("stage", "lot_amount", "complaint_deadline"):
            self.assertIn(c, rej)                           # v4
        for c in ("contact_region", "contact_address"):
            self.assertIn(c, ct)                            # v4

    def test_v2_duplicates_are_cancelled_by_migration(self):
        """У базі v2 компанія могла мати два тестові листи. Міграція лишає один."""
        import sqlite3
        now, to = E.now_iso(), E.TEST_RECIPIENTS[0]
        con = sqlite3.connect(self.db_path)
        con.executescript(self._schema_v1_sql())
        con.execute("INSERT INTO schema_meta(key,value) VALUES ('version','1')")
        con.execute("ALTER TABLE outreach_events ADD COLUMN origin TEXT NOT NULL"
                    " DEFAULT 'ENGINE'")
        con.execute("UPDATE schema_meta SET value='2' WHERE key='version'")
        con.execute("INSERT INTO companies VALUES (?,?,?,?)",
                    ("cid", "37406974", "ТОВ Двічі", now))
        con.execute("INSERT INTO test_recipient_allowlist VALUES (?,?,?)",
                    (to, now, None))
        for n, at in ((1, "2026-08-01T10:00:00+03:00"),
                      (2, "2026-08-02T10:00:00+03:00")):
            con.execute(
                "INSERT INTO outreach_events(outreach_id, company_id,"
                " outreach_type, mode, status, subject_rendered, body_rendered,"
                " delivery_email_actual, message_id_header, generated_at)"
                " VALUES (?,?,'FIRST_TOUCH','TEST','SENT','с','т',?,?,?)",
                (f"o{n}", "cid", to, f"<m{n}@tenderwin.com.ua>", at))
        con.commit()
        con.close()
        db = E.Db(self.db_path)                     # тут працює міграція 2→3
        rows = {r["outreach_id"]: (r["status"], r["error_code"]) for r in
                db.q("SELECT outreach_id, status, error_code FROM outreach_events")}
        db.close()
        self.assertEqual(rows["o1"][0], "SENT", "найраніший лист лишається")
        self.assertEqual(rows["o2"], ("CANCELLED", "DUPLICATE_FIRST_TOUCH"))

    def test_audit_log_records_the_whole_path(self):
        eng = self.engine([], mode="LIVE")
        self.run_cycle(eng, [facts()])
        eng.send()
        events = [r["event"] for r in eng.db.q("SELECT event FROM audit_log")]
        eng.close()
        for expected in ("LEAD_DISCOVERED", "OUTREACH_RESERVED",
                         "GMAIL_SEND_STARTED", "GMAIL_SEND_CONFIRMED"):
            self.assertIn(expected, events)


class TestVariantBOnly(Base):
    """
    Рішення власника продукту: усі листи — коротким варіантом B.

    Д-37: з вимкненим A/B рядка експерименту в базі немає, і зовнішній ключ
    падав на кожному листі. Лід ставав NOT_ELIGIBLE, тобто розсилка мовчки
    зупинялась, а звіт видавав це за штатну дедуплікацію.
    """

    def _eng(self, enabled: bool, db=None):
        cfg = E.Config(db_path=db or self.db_path,
                       out_dir=os.path.join(self.tmp, "out"),
                       mode="TEST", min_delay=0, max_delay=0, verbose=False,
                       experiment_enabled=enabled)
        eng = E.LeadEngine(cfg, E.FixtureSource([]),
                           E.DryRunTransport(os.path.join(self.tmp, "eml")), None)
        eng.repo.mark_history_migrated("тест")
        return eng

    def test_letters_still_go_out_with_the_experiment_off(self):
        eng = self._eng(False)
        q = self.run_cycle(eng, [facts()])
        reason = eng.db.one("SELECT ineligible_reason FROM leads")
        eng.close()
        self.assertEqual(q["QUEUED"], 1,
                         f"лист має ставитись у чергу; причина: {dict(reason)}")

    def test_everyone_gets_variant_b(self):
        eng = self._eng(False)
        self.run_cycle(eng, [facts(edrpou="11111111"),
                             facts(edrpou="22222222", ua="UA-2026-08-28-000002-a",
                                   bid="b2", award="a2", email="b@x.ua")])
        variants = {r["variant"] for r in
                    eng.db.q("SELECT variant FROM outreach_events")}
        eng.close()
        self.assertEqual(variants, {"B"}, "варіант A більше не надсилається")

    def test_no_dangling_experiment_reference(self):
        """Порожнє посилання краще за посилання на неіснуючий рядок."""
        eng = self._eng(False)
        self.run_cycle(eng, [facts()])
        row = eng.db.one("SELECT experiment_id FROM outreach_events")
        eng.close()
        self.assertIsNone(row["experiment_id"])

    def test_variant_is_assigned_when_the_experiment_is_on(self):
        """A/B не видалено — його можна ввімкнути назад одним значенням."""
        eng = self._eng(True)
        self.run_cycle(eng, [facts()])
        events = [r["event"] for r in eng.db.q("SELECT event FROM audit_log")]
        eng.close()
        self.assertIn("VARIANT_ASSIGNED", events)

    def test_integrity_error_is_not_reported_as_deduplication(self):
        """Д-37: будь-яка помилка цілісності видавалась за «вже писали»."""
        self.assertIn("QUEUE_INSERT_FAILED", E.ERR)
        err = E.ERR["QUEUE_INSERT_FAILED"]("FOREIGN KEY constraint failed")
        self.assertIn("FOREIGN KEY", err.what)
        self.assertNotIn("вже отримувала", err.what)


class TestPacing(Base):
    """Сухий прогін не має спати між листами."""

    def test_dry_run_is_not_paced(self):
        eng = self.engine([], mode="TEST",
                          transport=E.DryRunTransport(os.path.join(self.tmp, "e")))
        eng.sender.min_delay, eng.sender.max_delay = 120, 180
        paced = eng.sender.paced
        eng.close()
        self.assertFalse(paced, "у сухому прогоні надсилати нікуди — пауза зайва")

    def test_real_transport_is_paced(self):
        eng = self.engine([], mode="LIVE",
                          transport=FakeGmailTransport(os.path.join(self.tmp, "e")))
        eng.sender.min_delay, eng.sender.max_delay = 120, 180
        paced = eng.sender.paced
        eng.close()
        self.assertTrue(paced)

    def test_zero_delay_disables_pacing(self):
        eng = self.engine([], mode="LIVE",
                          transport=FakeGmailTransport(os.path.join(self.tmp, "e")))
        eng.sender.min_delay = eng.sender.max_delay = 0
        paced = eng.sender.paced
        eng.close()
        self.assertFalse(paced)

    def test_dry_run_of_many_leads_is_fast(self):
        import time as _t
        leads = [facts(edrpou=f"1000000{i}", ua=f"UA-2026-08-28-0000{i}-a",
                       bid=f"b{i}", award=f"a{i}") for i in range(5)]
        eng = self.engine([], mode="TEST",
                          transport=E.DryRunTransport(os.path.join(self.tmp, "e")))
        eng.sender.min_delay, eng.sender.max_delay = 120, 180
        self.run_cycle(eng, leads)
        t0 = _t.time()
        stats = eng.send()
        took = _t.time() - t0
        eng.close()
        self.assertEqual(stats["SENT"], 5)
        self.assertLess(took, 5, "сухий прогін 5 листів не має тривати хвилини")


class TestEmptyQueueExplainsItself(Base):
    """Порожня черга має пояснювати себе, а не друкувати нулі."""

    def test_lead_summary_shows_statuses(self):
        eng = self.engine([], mode="LIVE")
        self.run_cycle(eng, [facts(), facts(edrpou="99999999", ua="UA-2-a",
                                            bid="b2", award="a2", reason="")])
        summary = eng.lead_summary()
        eng.close()
        self.assertIn("READY_FOR_OUTREACH", summary)
        self.assertIn("MANUAL_REVIEW_REQUIRED", summary)

    def test_why_not_queued_names_the_reason(self):
        eng = self.engine([], mode="LIVE")
        self.run_cycle(eng, [facts(reason="")])
        reasons = eng.why_not_queued()
        eng.close()
        self.assertTrue(reasons)
        self.assertIn("підстав", reasons[0]["reason"])

    def test_pending_summary_shows_leftover_queue(self):
        eng = self.engine([], mode="LIVE")
        self.run_cycle(eng, [facts()])
        pend = eng.pending_summary()
        eng.close()
        self.assertEqual(pend.get("QUEUED"), 1,
                         "лист, який ще не пішов, має бути видно")


class TestOneCommandFlow(Base):
    """Токен зберігається сам, секрет у Colab не обовʼязковий."""

    def setUp(self):
        super().setUp()
        self._orig_dir = E.TOKEN_DIR
        E.TOKEN_DIR = self.tmp

    def tearDown(self):
        E.TOKEN_DIR = self._orig_dir
        super().tearDown()

    def test_token_round_trip(self):
        E.save_token({"client_id": "a", "client_secret": "b",
                      "refresh_token": "c"}, self.tmp)
        raw = E.load_token()
        self.assertIsNotNone(raw)
        self.assertEqual(json.loads(raw)["refresh_token"], "c")

    def test_provider_falls_back_to_file(self):
        os.environ.pop("GMAIL_OAUTH_JSON", None)
        self.assertIsNone(E._provider_with_file("GMAIL_OAUTH_JSON"))
        E.save_token({"client_id": "a", "client_secret": "b",
                      "refresh_token": "c"}, self.tmp)
        got = E._provider_with_file("GMAIL_OAUTH_JSON")
        self.assertIn("refresh_token", got or "")

    def test_env_wins_over_file(self):
        E.save_token({"client_id": "file", "client_secret": "b",
                      "refresh_token": "c"}, self.tmp)
        os.environ["GMAIL_OAUTH_JSON"] = json.dumps(
            {"client_id": "env", "client_secret": "b", "refresh_token": "c"})
        try:
            got = json.loads(E._provider_with_file("GMAIL_OAUTH_JSON"))
            self.assertEqual(got["client_id"], "env")
        finally:
            os.environ.pop("GMAIL_OAUTH_JSON", None)

    def test_token_file_is_hidden_and_beside_the_database(self):
        path = E.save_token({"client_id": "a", "client_secret": "b",
                             "refresh_token": "c"}, self.tmp)
        self.assertTrue(os.path.basename(path).startswith("."))
        self.assertEqual(os.path.dirname(path), self.tmp)

    def test_setup_gmail_without_client_file_explains(self):
        import io, contextlib
        orig = E.find_oauth_client
        try:
            E.find_oauth_client = lambda dirs=None: []
            buf = io.StringIO()
            with contextlib.redirect_stdout(buf):
                ok = E.setup_gmail(self.tmp)
        finally:
            E.find_oauth_client = orig
        self.assertFalse(ok)
        self.assertIn("НЕ ЗНАЙДЕНО ФАЙЛ КЛІЄНТА OAUTH", buf.getvalue())
        self.assertIn("Desktop app", buf.getvalue())

    def test_setup_gmail_full_flow_with_mocks(self):
        """Проходимо весь шлях, підмінивши браузер і обмін коду."""
        import io, contextlib, builtins
        f = os.path.join(self.tmp, "client_secret_x.apps.googleusercontent.com.json")
        open(f, "w").write(json.dumps({"installed": {
            "client_id": "604019.apps.googleusercontent.com",
            "client_secret": "GOCSPX-x"}}))

        class _Resp:
            status_code = 200
            content = b"x"
            def json(self):
                return {"refresh_token": "1//REFRESH", "access_token": "a"}

        import requests
        orig_post, orig_input = requests.post, builtins.input
        orig_check = E.check_gmail_setup
        try:
            requests.post = lambda *a, **k: _Resp()
            builtins.input = lambda *a: "http://localhost/?code=4/0ABC"
            E.check_gmail_setup = lambda *a, **k: {"кроки": [{"ok": True}],
                                                   "підсумок": "ok"}
            buf = io.StringIO()
            with contextlib.redirect_stdout(buf):
                ok = E.setup_gmail(self.tmp, client_secret_file=f)
        finally:
            requests.post, builtins.input = orig_post, orig_input
            E.check_gmail_setup = orig_check

        self.assertTrue(ok, buf.getvalue()[-500:])
        saved = json.loads(E.load_token())
        self.assertEqual(saved["refresh_token"], "1//REFRESH")
        self.assertIn("Токен збережено", buf.getvalue())
        self.assertNotIn("1//REFRESH", buf.getvalue(),
                         "токен не має потрапляти у вивід ноутбука")
        self.assertNotIn("Colab Secrets", buf.getvalue(),
                         "не радимо створювати секрет, коли токен уже у файлі")

    def test_setup_gmail_skips_when_token_already_works(self):
        import io, contextlib
        E.save_token({"client_id": "a", "client_secret": "b",
                      "refresh_token": "c"}, self.tmp)
        orig = E.check_gmail_setup
        try:
            E.check_gmail_setup = lambda *a, **k: {"кроки": [{"ok": True}],
                                                   "підсумок": "ok"}
            with contextlib.redirect_stdout(io.StringIO()):
                ok = E.setup_gmail(self.tmp)
        finally:
            E.check_gmail_setup = orig
        self.assertTrue(ok)


class TestDryRunDoesNotConsumeQueue(Base):
    """
    Сухий прогін не має «витрачати» листи.

    Саме через це листи не йшли: сухий прогін позначав їх надісланими,
    черга порожніла, і реальна відправка не знаходила нічого.
    """

    def test_preview_leaves_letters_in_queue(self):
        eng = self.engine([], mode="TEST")
        self.run_cycle(eng, [facts()])
        before = eng.pending_summary()
        res = eng.preview()
        after = eng.pending_summary()
        eng.close()
        self.assertEqual(before.get("QUEUED"), 1)
        self.assertEqual(res["складено"], 1)
        self.assertEqual(res["надіслано"], 0)
        self.assertEqual(after.get("QUEUED"), 1,
                         "після сухого прогону лист має лишитись у черзі")

    def test_preview_writes_readable_files(self):
        eng = self.engine([], mode="TEST")
        self.run_cycle(eng, [facts()])
        eng.preview()
        d = os.path.join(eng.cfg.out_dir, "eml")
        files = sorted(os.listdir(d)) if os.path.isdir(d) else []
        eng.close()
        self.assertEqual(len(files), 2, "поруч з .eml має бути читабельний .txt")
        txt = [f for f in files if f.endswith(".txt")][0]
        with open(os.path.join(d, txt), encoding="utf-8") as fh:
            body = fh.read()
        # Без підключеного лід-генератора кличний відмінок недоступний,
        # тому тут безпечне звертання — саме так і має бути (§7).
        self.assertIn("Добрий день!", body)
        self.assertIn("не надано довідку", body)
        self.assertIn("Кому:", body)
        self.assertIn("Тема:", body)

    def test_preview_eml_is_valid_mime(self):
        """.eml має відкриватись поштовою програмою."""
        import email, email.policy
        eng = self.engine([], mode="TEST")
        self.run_cycle(eng, [facts()])
        eng.preview()
        d = os.path.join(eng.cfg.out_dir, "eml")
        eml = [f for f in os.listdir(d) if f.endswith(".eml")][0]
        eng.close()
        with open(os.path.join(d, eml), "rb") as fh:
            msg = email.message_from_bytes(fh.read(), policy=email.policy.default)
        self.assertIn("Добрий день!", msg.get_content())
        self.assertIn("не надано довідку", msg.get_content())
        self.assertIn("@", msg["To"])
        self.assertIn("Щодо закупівлі", msg["Subject"])

    def test_preview_file_name_shows_the_real_recipient(self):
        eng = self.engine([], mode="TEST")
        self.run_cycle(eng, [facts(email="klient@example.com")])
        eng.preview()
        d = os.path.join(eng.cfg.out_dir, "eml")
        names = os.listdir(d)
        eng.close()
        self.assertTrue(any("klient" in n for n in names),
                        f"за назвою файлу має бути видно, кому лист: {names}")

    def test_preview_never_uses_the_transport(self):
        """Надіслати із сухого прогону має бути структурно неможливо."""
        class Explode:
            name = "GMAIL"
            def send(self, *a, **k):
                raise AssertionError("сухий прогін не сміє надсилати")
            def find_by_message_id(self, *a, **k):
                return None
        eng = self.engine([], mode="TEST", transport=Explode())
        self.run_cycle(eng, [facts()])
        res = eng.preview()
        eng.close()
        self.assertEqual(res["складено"], 1)

    def test_dry_then_live_actually_sends(self):
        """Повний шлях: спершу подивились, потім надіслали."""
        eng = self.engine([], mode="TEST")
        self.run_cycle(eng, [facts()])
        eng.preview()
        eng.close()

        eng2 = self.engine([], mode="TEST")
        stats = eng2.send()
        row = eng2.db.one("SELECT status, gmail_message_id FROM outreach_events")
        eng2.close()
        self.assertEqual(stats["SENT"], 1, "після сухого прогону лист має піти")
        self.assertEqual(row["status"], "SENT")

    def test_reset_repairs_old_broken_database(self):
        """Ремонт баз, які вже зіпсував старий сухий прогін."""
        eng = self.engine([], mode="TEST")
        self.run_cycle(eng, [facts()])
        eng.send()                      # старий сценарій: DryRun позначив SENT
        self.assertEqual(eng.pending_summary().get("SENT"), 1)
        eng.close()

        rep = E.reset_dry_run_sends(self.db_path)
        self.assertEqual(rep["повернуто_в_чергу"], 1)

        eng2 = self.engine([], mode="TEST")
        pend = eng2.pending_summary()
        eng2.close()
        self.assertEqual(pend.get("QUEUED"), 1)

    def test_reset_does_not_touch_real_sends(self):
        eng = self.engine([], mode="LIVE",
                          transport=FakeGmailTransport(os.path.join(self.tmp, "e")))
        self.run_cycle(eng, [facts()])
        eng.send()
        oid = eng.db.one("SELECT outreach_id FROM outreach_events")["outreach_id"]
        eng.db.conn.execute(
            "UPDATE outreach_events SET gmail_message_id='real-123' WHERE outreach_id=?",
            (oid,))
        eng.close()
        rep = E.reset_dry_run_sends(self.db_path)
        self.assertEqual(rep["повернуто_в_чергу"], 0,
                         "справжні відправки чіпати не можна")

    def test_why_no_letters_names_the_queue(self):
        import io, contextlib
        eng = self.engine([], mode="TEST")
        self.run_cycle(eng, [facts()])
        eng.preview()
        eng.close()
        buf = io.StringIO()
        with contextlib.redirect_stdout(buf):
            E.why_no_letters(self.db_path)
        out = buf.getvalue()
        self.assertIn("чекають у черзі", out)
        self.assertIn("E.go(live=True)", out)

    def test_why_no_letters_detects_broken_database(self):
        import io, contextlib
        eng = self.engine([], mode="TEST")
        self.run_cycle(eng, [facts()])
        eng.send()                      # старий зламаний сценарій
        eng.close()
        buf = io.StringIO()
        with contextlib.redirect_stdout(buf):
            E.why_no_letters(self.db_path)
        out = buf.getvalue()
        self.assertIn("СУХИМ прогоном", out)
        self.assertIn("reset_dry_run_sends", out)


class TestStorageGate2(Base):
    """
    Тихо працювати з непідключеним Диском не можна.

    Ці тести НЕ дивляться на справжній Диск і ніколи його не монтують:
    інакше вони падали б на машині, де Диск підключений, і посеред
    прогону перемонтовували б людині сховище.
    """

    def setUp(self):
        super().setUp()
        self._orig = (E.DRIVE_PREFIX, E.DRIVE_MOUNT_ROOT)
        # «Диск», якого свідомо немає
        E.DRIVE_PREFIX = os.path.join(self.tmp, "fakedrive")
        E.DRIVE_MOUNT_ROOT = os.path.join(self.tmp, "fakedrive", "MyDrive")
        self.fake = os.path.join(E.DRIVE_MOUNT_ROOT, "TENDERWIN")

    def tearDown(self):
        E.DRIVE_PREFIX, E.DRIVE_MOUNT_ROOT = self._orig
        super().tearDown()

    def test_unmounted_drive_path_is_detected(self):
        self.assertFalse(E.drive_is_mounted(self.fake))

    def test_mounted_drive_path_is_accepted(self):
        orig = os.path.ismount
        try:
            os.path.ismount = lambda p: p == E.DRIVE_PREFIX
            self.assertTrue(E.drive_is_mounted(self.fake))
        finally:
            os.path.ismount = orig

    def test_fake_folders_do_not_fool_the_check(self):
        """
        Створені локально папки з тими самими назвами — не Диск.
        Саме так ламалась перша версія перевірки: os.makedirs створює
        весь ланцюжок, включно з MyDrive, і перевірка на існування
        каталогу після цього завжди казала «усе гаразд».
        """
        os.makedirs(self.fake, exist_ok=True)
        self.assertTrue(os.path.isdir(E.DRIVE_MOUNT_ROOT))
        self.assertTrue(os.path.isdir(self.fake))
        self.assertFalse(E.drive_is_mounted(self.fake),
                         "каталог не є точкою монтування")

    def test_local_path_is_always_considered_available(self):
        self.assertTrue(E.drive_is_mounted(self.tmp))
        self.assertTrue(E.drive_is_mounted("/content/TENDERWIN"))

    def test_go_refuses_unmounted_drive(self):
        import io, contextlib
        buf = io.StringIO()
        with contextlib.redirect_stdout(buf):
            r = E.go(live=False, dysk=self.fake)
        out = buf.getvalue()
        self.assertIsNone(r, "прогін не має починатись без сховища")
        self.assertIn("ДИСК НЕ ПІДКЛЮЧЕНО", out)
        self.assertFalse(os.path.isdir(self.fake),
                         "не створювати підробну папку замість Диска")

    def test_setup_gmail_refuses_unmounted_drive(self):
        import io, contextlib
        buf = io.StringIO()
        with contextlib.redirect_stdout(buf):
            ok = E.setup_gmail(self.fake)
        self.assertFalse(ok)
        self.assertIn("ДИСК НЕ ПІДКЛЮЧЕНО", buf.getvalue())

    def test_gate_never_remounts_real_drive(self):
        """Тест не має права чіпати сховище людини."""
        import io, contextlib
        calls = []
        orig = E.ensure_drive
        try:
            E.ensure_drive = lambda p, try_remount=True: (calls.append(p), False)[1]
            with contextlib.redirect_stdout(io.StringIO()):
                E.go(live=False, dysk=self.fake)
        finally:
            E.ensure_drive = orig
        self.assertEqual(calls, [self.fake])

    def test_local_fallback_warns_about_losing_state(self):
        """
        run_colab підмінено навмисно: без цього тест пішов би у справжню
        мережу до Prozorro — повільно, ненадійно і без потреби.
        """
        import io, contextlib
        local = os.path.join(self.tmp, "TENDERWIN")
        orig = E.run_colab
        try:
            E.run_colab = lambda **kw: {"stub": True}
            buf = io.StringIO()
            with contextlib.redirect_stdout(buf):
                r = E.go(live=False, dysk=local)
        finally:
            E.run_colab = orig
        self.assertEqual(r, {"stub": True})
        self.assertIn("зникне разом із сеансом", buf.getvalue())

    def test_go_never_touches_network_before_storage_check(self):
        """Порядок важливий: сховище перевіряється ПЕРЕД будь-якою роботою."""
        orig = E.run_colab
        called = []
        try:
            E.run_colab = lambda **kw: called.append(kw)
            import io, contextlib
            with contextlib.redirect_stdout(io.StringIO()):
                E.go(live=False, dysk=self.fake)
        finally:
            E.run_colab = orig
        self.assertEqual(called, [], "прогін не мав початись")


class TestPastedUrlHelp(Base):
    """Найчастіша помилка при копіюванні — неповна адреса."""

    def test_bare_localhost_is_recognised(self):
        with self.assertRaises(ValueError) as ctx:
            E.extract_code("localhost")
        self.assertIn("лише «localhost»", str(ctx.exception))

    def test_url_without_code_explains_how_to_copy(self):
        with self.assertRaises(ValueError) as ctx:
            E.extract_code("http://localhost/")
        self.assertIn("Ctrl+L", str(ctx.exception))

    def test_access_denied_explains_test_users(self):
        with self.assertRaises(ValueError) as ctx:
            E.extract_code("http://localhost/?error=access_denied")
        self.assertIn("Test users", str(ctx.exception))

    def test_full_url_still_works(self):
        self.assertEqual(
            E.extract_code("http://localhost/?code=4%2F0AX4-TEST&scope=x"),
            "4/0AX4-TEST")

    def test_setup_gmail_retries_on_bad_paste(self):
        """Погана вставка не має вимагати починати все спочатку."""
        import io, contextlib, builtins, json as _j
        f = os.path.join(self.tmp, "client_secret_a.apps.googleusercontent.com.json")
        open(f, "w").write(_j.dumps({"installed": {
            "client_id": "604019.apps.googleusercontent.com",
            "client_secret": "GOCSPX-x"}}))

        class _Resp:
            status_code, content = 200, b"x"
            def json(self): return {"refresh_token": "1//OK"}

        import requests
        answers = iter(["localhost", "http://localhost/?code=4/0GOOD"])
        orig = (requests.post, builtins.input, E.check_gmail_setup)
        try:
            requests.post = lambda *a, **k: _Resp()
            builtins.input = lambda *a: next(answers)
            E.check_gmail_setup = lambda *a, **k: {"кроки": [{"ok": True}],
                                                   "підсумок": "ok"}
            buf = io.StringIO()
            with contextlib.redirect_stdout(buf):
                ok = E.setup_gmail(self.tmp, client_secret_file=f)
        finally:
            requests.post, builtins.input, E.check_gmail_setup = orig
        self.assertTrue(ok, buf.getvalue()[-400:])
        self.assertIn("Спробуйте ще раз", buf.getvalue())
        self.assertEqual(_j.loads(E.load_token())["refresh_token"], "1//OK")


class TestOAuthHelpers(Base):
    """Те, що можна перевірити без мережі: посилання і розбір відповіді."""

    def test_auth_url_asks_for_offline_access(self):
        url = E.build_auth_url("123.apps.googleusercontent.com")
        self.assertIn("access_type=offline", url,
                      "без offline Google не видасть refresh_token")
        self.assertIn("prompt=consent", url)
        self.assertIn("gmail.send", url)
        self.assertIn("gmail.readonly", url)

    def test_extract_code_from_pasted_url(self):
        u = "http://localhost/?code=4%2F0AX4XfWh-TEST&scope=https%3A%2F%2Fmail"
        self.assertEqual(E.extract_code(u), "4/0AX4XfWh-TEST")

    def test_extract_code_from_bare_code(self):
        self.assertEqual(E.extract_code("4/0AX4XfWh-TEST"), "4/0AX4XfWh-TEST")

    def test_extract_code_reports_google_error(self):
        with self.assertRaises(ValueError) as ctx:
            E.extract_code("http://localhost/?error=access_denied")
        self.assertIn("access_denied", str(ctx.exception))

    def test_extract_code_rejects_url_without_code(self):
        with self.assertRaises(ValueError):
            E.extract_code("http://localhost/")

    def test_extract_code_rejects_empty(self):
        with self.assertRaises(ValueError):
            E.extract_code("   ")

    def test_missing_secret_is_typed_not_crash(self):
        old = os.environ.pop("GMAIL_OAUTH_JSON", None)
        try:
            with self.assertRaises(E.EngineError) as ctx:
                E._oauth_token_info()
            self.assertEqual(ctx.exception.code, "TRANSPORT_UNAVAILABLE")
        finally:
            if old:
                os.environ["GMAIL_OAUTH_JSON"] = old

    def test_incomplete_token_names_the_missing_fields(self):
        os.environ["GMAIL_OAUTH_JSON"] = '{"client_id": "x"}'
        try:
            with self.assertRaises(E.EngineError) as ctx:
                E._oauth_token_info()
            self.assertIn("client_secret", ctx.exception.what)
            self.assertIn("refresh_token", ctx.exception.what)
        finally:
            os.environ.pop("GMAIL_OAUTH_JSON", None)

    def test_broken_json_is_reported_as_json_problem(self):
        os.environ["GMAIL_OAUTH_JSON"] = "не json"
        try:
            with self.assertRaises(E.EngineError) as ctx:
                E._oauth_token_info()
            self.assertIn("JSON", ctx.exception.what)
        finally:
            os.environ.pop("GMAIL_OAUTH_JSON", None)

    def test_valid_token_info_passes_validation(self):
        os.environ["GMAIL_OAUTH_JSON"] = json.dumps(
            {"client_id": "a", "client_secret": "b", "refresh_token": "c"})
        try:
            info = E._oauth_token_info()
            self.assertEqual(info["refresh_token"], "c")
        finally:
            os.environ.pop("GMAIL_OAUTH_JSON", None)

    def test_transport_selection_prefers_oauth(self):
        """Підміняємо джерело секретів — тест не залежить від машини."""
        cfg = E.Config(db_path=self.db_path, out_dir=self.tmp, mode="LIVE")
        self.assertEqual(E.make_transport(cfg, dry_run=True).name, "DRY_RUN")
        orig = E.SECRET_PROVIDER
        try:
            E.SECRET_PROVIDER = lambda name: None
            with self.assertRaises(E.EngineError) as ctx:
                E.make_transport(cfg, dry_run=False)
            self.assertIn("GMAIL_OAUTH_JSON", ctx.exception.what)
        finally:
            E.SECRET_PROVIDER = orig

    def test_diagnostic_without_any_secret_points_to_oauth(self):
        """
        Коли нічого не налаштовано, діагностика має вести до OAuth,
        а не в глухий кут зі службовим акаунтом: у Workspace ключі
        здебільшого заборонені за замовчуванням.
        """
        import io, contextlib
        orig = E.SECRET_PROVIDER
        try:
            E.SECRET_PROVIDER = lambda name: None
            buf = io.StringIO()
            with contextlib.redirect_stdout(buf):
                rep = E.check_gmail_setup(sender_email="vitalii@tenderwin.com.ua",
                                          send_probe=False)
            out = buf.getvalue()
        finally:
            E.SECRET_PROVIDER = orig
        # без бібліотек Google діагностика зупиняється раніше — це теж коректно
        if "google-api-python-client" in rep["підсумок"]:
            self.skipTest("бібліотеки Google не встановлені в цьому середовищі")
        self.assertIn("GMAIL_OAUTH_JSON", rep["підсумок"])
        self.assertIn("Desktop app", out)
        self.assertNotIn("Немає секрету GMAIL_SA_JSON", rep["підсумок"])

    def test_identifies_service_account_key(self):
        import io, contextlib, json as _j
        f = os.path.join(self.tmp, "sa.json")
        open(f, "w").write(_j.dumps({
            "type": "service_account", "project_id": "p",
            "client_email": "bot@p.iam.gserviceaccount.com", "client_id": "1",
            "private_key": "-----BEGIN PRIVATE KEY-----"}))
        buf = io.StringIO()
        with contextlib.redirect_stdout(buf):
            r = E.whats_this_file(f)
        self.assertEqual(r["вид"], "SERVICE_ACCOUNT_KEY")
        self.assertFalse(r["придатний"])
        self.assertIn("НЕ потрібен", buf.getvalue())

    def test_identifies_oauth_desktop_client(self):
        import io, contextlib, json as _j
        f = os.path.join(self.tmp, "client_secret_x.json")
        open(f, "w").write(_j.dumps({
            "installed": {"client_id": "123.apps.googleusercontent.com",
                          "client_secret": "GOCSPX-x"}}))
        with contextlib.redirect_stdout(io.StringIO()):
            r = E.whats_this_file(f)
        self.assertEqual(r["вид"], "OAUTH_CLIENT")
        self.assertTrue(r["придатний"])
        self.assertEqual(r["client_secret"], "GOCSPX-x")

    def test_identifies_our_own_token(self):
        import io, contextlib, json as _j
        f = os.path.join(self.tmp, "t.json")
        open(f, "w").write(_j.dumps({"client_id": "a", "client_secret": "b",
                                     "refresh_token": "c"}))
        with contextlib.redirect_stdout(io.StringIO()):
            r = E.whats_this_file(f)
        self.assertEqual(r["вид"], "TENDERWIN_TOKEN")

    def test_missing_path_says_not_found_not_bad_json(self):
        """
        Саме тут людина застрягла: у змінній лишився приклад шляху,
        а повідомлення казало «це не JSON» і збивало зі сліду.
        """
        import io, contextlib
        buf = io.StringIO()
        with contextlib.redirect_stdout(buf):
            r = E.whats_this_file("/content/client_secret_....json")
        self.assertEqual(r["вид"], "FILE_NOT_FOUND")
        self.assertIn("ФАЙЛ НЕ ЗНАЙДЕНО", buf.getvalue())
        self.assertNotIn("Це не JSON", buf.getvalue())

    def test_finds_client_file_by_name(self):
        import json as _j
        name = ("client_secret_604019504962-abc"
                ".apps.googleusercontent.com.json")
        open(os.path.join(self.tmp, name), "w").write(_j.dumps(
            {"installed": {"client_id": "x", "client_secret": "y"}}))
        found = E.find_oauth_client([self.tmp])
        self.assertEqual(len(found), 1)
        self.assertTrue(found[0].endswith(name))

    def test_mint_without_file_gives_actionable_error(self):
        orig = E.find_oauth_client
        try:
            E.find_oauth_client = lambda dirs=None: []
            with self.assertRaises(E.EngineError) as ctx:
                E.mint_oauth_token()
            self.assertEqual(ctx.exception.code, "OAUTH_CLIENT_FILE_NOT_FOUND")
            self.assertIn("завантаження", ctx.exception.next_step)
        finally:
            E.find_oauth_client = orig

    def test_mint_autofinds_single_client_file(self):
        import io, contextlib, json as _j
        f = os.path.join(self.tmp, "client_secret_1.apps.googleusercontent.com.json")
        open(f, "w").write(_j.dumps({"installed": {
            "client_id": "604019.apps.googleusercontent.com",
            "client_secret": "GOCSPX-x"}}))
        orig = E.find_oauth_client
        try:
            E.find_oauth_client = lambda dirs=None: [f]
            buf = io.StringIO()
            with contextlib.redirect_stdout(buf):
                res = E.mint_oauth_token()
            out = buf.getvalue()
        finally:
            E.find_oauth_client = orig
        self.assertIsNone(res)
        self.assertIn("Знайдено файл клієнта", out)
        self.assertIn("accounts.google.com", out)
        self.assertIn("604019", out)

    def test_identifies_garbage(self):
        import io, contextlib
        f = os.path.join(self.tmp, "junk.json")
        open(f, "w").write("не json зовсім")
        with contextlib.redirect_stdout(io.StringIO()):
            r = E.whats_this_file(f)
        self.assertEqual(r["вид"], "NOT_JSON")

    def test_mint_refuses_service_account_key(self):
        import io, contextlib, json as _j
        f = os.path.join(self.tmp, "sa2.json")
        open(f, "w").write(_j.dumps({"type": "service_account",
                                     "client_email": "b@p", "client_id": "1"}))
        with contextlib.redirect_stdout(io.StringIO()):
            with self.assertRaises(E.EngineError) as ctx:
                E.mint_oauth_token(client_secret_file=f)
        self.assertEqual(ctx.exception.code, "WRONG_CREDENTIAL_TYPE")

    def test_mint_accepts_oauth_client_file(self):
        import io, contextlib, json as _j
        f = os.path.join(self.tmp, "c.json")
        open(f, "w").write(_j.dumps({
            "installed": {"client_id": "123.apps.googleusercontent.com",
                          "client_secret": "GOCSPX-x"}}))
        buf = io.StringIO()
        with contextlib.redirect_stdout(buf):
            res = E.mint_oauth_token(client_secret_file=f)   # крок 1
        self.assertIsNone(res, "крок 1 лише друкує посилання")
        self.assertIn("accounts.google.com", buf.getvalue())
        self.assertIn("access_type=offline", buf.getvalue())

    def test_mint_without_anything_says_what_to_do(self):
        """Без файлу і без значень — помилка має бути дієвою, не абстрактною."""
        orig = E.find_oauth_client
        try:
            E.find_oauth_client = lambda dirs=None: []
            with self.assertRaises(E.EngineError) as ctx:
                E.mint_oauth_token()
        finally:
            E.find_oauth_client = orig
        self.assertEqual(ctx.exception.code, "OAUTH_CLIENT_FILE_NOT_FOUND")

    def test_mint_with_id_but_no_secret_is_reported(self):
        with self.assertRaises(E.EngineError) as ctx:
            E.mint_oauth_token(client_id="123.apps.googleusercontent.com")
        self.assertEqual(ctx.exception.code, "MISSING_OAUTH_CLIENT")

    def test_colab_secret_does_not_raise_without_colab(self):
        """
        Обгортка має повертати None, а не падати. Саме на цьому падала
        клітинка перевірки: userdata.get() кидає SecretNotFoundError
        рівно тоді, коли перевірка й потрібна.
        """
        self.assertIsNone(E.colab_secret("GMAIL_OAUTH_JSON_НЕМАЄ_ТАКОГО"))

    def test_colab_state_reads_empty_database(self):
        eng = self.engine([], mode="TEST")
        eng.close()
        import io, contextlib
        buf = io.StringIO()
        with contextlib.redirect_stdout(buf):
            E.colab_state(self.db_path, self.tmp)
        out = buf.getvalue()
        self.assertIn("ЛІДИ", out)
        self.assertIn("ЗВЕРНЕННЯ", out)

    def test_colab_state_shows_real_leads(self):
        eng = self.engine([], mode="LIVE")
        self.run_cycle(eng, [facts()])
        eng.close()
        import io, contextlib
        buf = io.StringIO()
        with contextlib.redirect_stdout(buf):
            E.colab_state(self.db_path, self.tmp)
        out = buf.getvalue()
        self.assertIn("READY_FOR_OUTREACH", out)
        self.assertIn("UA-2026-08-28-000001-a", out)

    def test_service_account_env_name_is_still_supported(self):
        """Запасний шлях лишається для тих, у кого ключ уже є."""
        self.assertEqual(E.GMAIL_SA_ENV, "GMAIL_SA_JSON")
        self.assertEqual(E.GMAIL_OAUTH_ENV, "GMAIL_OAUTH_JSON")

    def test_tests_never_read_real_secrets(self):
        """Захист від повернення дефекту: тести не лізуть у Colab Secrets."""
        self.assertFalse(E.ALLOW_COLAB_SECRETS)
        os.environ.pop("GMAIL_OAUTH_JSON", None)
        self.assertIsNone(E._default_secret_provider("GMAIL_OAUTH_JSON"))


if __name__ == "__main__":
    unittest.main(verbosity=2)


class TestNoDuplicateForTwoRejections(Base):
    """
    Вимога власника: «щоб не було дублювання листів учаснику з одним ЄДРПОУ,
    навіть якщо в нього відхилили дві пропозиції».
    """

    TWO = [facts(ua="UA-2026-08-07-008904-a", bid="b1", award="a1"),
           facts(ua="UA-2026-08-07-009602-a", bid="b2", award="a2")]

    def _queued(self, mode):
        eng = self.engine([], mode=mode)
        q = self.run_cycle(eng, self.TWO)
        rows = eng.db.q("SELECT status FROM outreach_events")
        eng.close()
        return q["QUEUED"], len(rows)

    def test_one_letter_in_live_mode(self):
        queued, rows = self._queued("LIVE")
        self.assertEqual(queued, 1, "дві пропозиції однієї фірми — один лист")
        self.assertEqual(rows, 1, "другий лист не має навіть створюватись")

    def test_one_letter_in_test_mode(self):
        queued, rows = self._queued("TEST")
        self.assertEqual(queued, 1,
                         "у тестовому режимі теж один лист на ЄДРПОУ")
        self.assertEqual(rows, 1)

    def test_second_rejection_is_still_recorded(self):
        """Лист один, але саме відхилення з бази зникати не має (правило 14)."""
        eng = self.engine([], mode="TEST")
        self.run_cycle(eng, self.TWO)
        rej = eng.db.q("SELECT rejection_id FROM rejections")
        comp = eng.db.q("SELECT company_id FROM companies")
        eng.close()
        self.assertEqual(len(rej), 2, "обидва відхилення лишаються в базі")
        self.assertEqual(len(comp), 1, "але компанія одна")

    def test_two_different_companies_get_two_letters(self):
        """Дедуплікація не має схлопувати РІЗНІ фірми."""
        eng = self.engine([], mode="TEST")
        q = self.run_cycle(eng, [
            facts(edrpou="12345678", ua="UA-2026-08-07-008904-a"),
            facts(edrpou="37406974", ua="UA-2026-08-07-008904-a",
                  bid="b2", award="a2", email="b@x.ua")])
        eng.close()
        self.assertEqual(q["QUEUED"], 2)

    def test_database_blocks_duplicate_even_if_code_tries(self):
        """
        §21: дублі блокує БАЗА, а не застереження в коді. Обходимо політику
        і вставляємо другий FIRST_TOUCH напряму — індекс має відмовити.
        """
        import sqlite3
        eng = self.engine([], mode="TEST")
        self.run_cycle(eng, self.TWO)
        row = eng.db.one("SELECT * FROM outreach_events LIMIT 1")
        with self.assertRaises(sqlite3.IntegrityError):
            eng.db.conn.execute(
                "INSERT INTO outreach_events(outreach_id, company_id,"
                " outreach_type, mode, status, subject_rendered, body_rendered,"
                " delivery_email_actual, message_id_header, generated_at)"
                " VALUES (?,?,'FIRST_TOUCH',?,'QUEUED','с','т',?,?,?)",
                ("dup", row["company_id"], row["mode"],
                 row["delivery_email_actual"], "<dup@tenderwin.com.ua>",
                 E.now_iso()))
        eng.close()


class TestCaseCards(Base):
    """Теки на Диску з ІД закупівлі та службова картка в кожній."""

    def _build(self, items=None, **kw):
        eng = self.engine([], mode="TEST")
        self.run_cycle(eng, items or [facts()])
        eng.send()
        eng.close()
        root = os.path.join(self.tmp, "drive")
        stats = E.build_case_cards(self.db_path, root, legacy=None, **kw)
        return root, stats

    def _text(self, path):
        import docx
        d = docx.Document(path)
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)

    def test_folder_is_named_by_tender_id(self):
        root, stats = self._build()
        folder = os.path.join(root, "spravy", "UA-2026-08-28-000001-a")
        self.assertTrue(os.path.isdir(folder), "тека має зватися ІД закупівлі")
        self.assertEqual(stats["створено"], 1)
        self.assertTrue(os.path.isfile(
            os.path.join(folder, "00_KARTKA_12345678.docx")))

    def test_card_starts_with_a_link_to_the_tender(self):
        root, _ = self._build()
        p = os.path.join(root, "spravy", "UA-2026-08-28-000001-a",
                         "00_KARTKA_12345678.docx")
        import docx
        d = docx.Document(p)
        xml = d.element.xml
        self.assertIn("prozorro.gov.ua/tender/UA-2026-08-28-000001-a", "".join(
            r.target_ref for r in d.part.rels.values() if r.is_external),
            "посилання має вести на цю саму закупівлю")
        i = xml.index("<w:hyperlink")
        j = xml.index("</w:hyperlink>")
        self.assertIn("UA-2026-08-28-000001-a", xml[i:j],
                      "видимий текст посилання — ІД закупівлі")
        self.assertLess(i, xml.index("<w:tbl"),
                        "посилання зверху, до таблиці")
        head = " ".join(p.text for p in d.paragraphs[:3])
        self.assertIn("UA-2026-08-28-000001-a", head,
                      "ІД закупівлі — у перших рядках картки")

    def test_card_keeps_the_verbatim_reason(self):
        root, _ = self._build()
        text = self._text(os.path.join(root, "spravy",
                                       "UA-2026-08-28-000001-a",
                                       "00_KARTKA_12345678.docx"))
        self.assertIn("довідку про наявність працівників", text,
                      "формулювання замовника переноситься дослівно")

    def test_unknown_fields_are_marked_not_invented(self):
        root, _ = self._build()
        text = self._text(os.path.join(root, "spravy",
                                       "UA-2026-08-28-000001-a",
                                       "00_KARTKA_12345678.docx"))
        self.assertIn("НЕ ВСТАНОВЛЕНО", text.upper(),
                      "невідоме поле має бути позначене, а не вигадане")

    def test_human_notes_are_never_overwritten(self):
        """Правило 36: людські примітки AI не перезаписує."""
        root, _ = self._build()
        p = os.path.join(root, "spravy", "UA-2026-08-28-000001-a",
                         "00_KARTKA_12345678.docx")
        import docx
        d = docx.Document(p)
        d.add_paragraph("ВИСНОВОК ЮРИСТА: оскаржувати недоцільно.")
        d.save(p)
        stats = E.build_case_cards(self.db_path, root, legacy=None)
        self.assertEqual(stats["створено"], 0)
        self.assertEqual(stats["пропущено"], 1)
        self.assertIn("ВИСНОВОК ЮРИСТА", self._text(p))

    def test_two_rejections_one_company_two_cards_one_letter(self):
        """Карток дві (по закупівлях), лист один (по ЄДРПОУ)."""
        two = TestNoDuplicateForTwoRejections.TWO
        root, stats = self._build(items=two)
        self.assertEqual(stats["закупівель"], 2)
        self.assertEqual(stats["створено"], 2)
        for ua in ("UA-2026-08-07-008904-a", "UA-2026-08-07-009602-a"):
            self.assertTrue(os.path.isfile(os.path.join(
                root, "spravy", ua, "00_KARTKA_12345678.docx")))


class TestCardNumbers(Base):
    """Дві плати — дві різні бази. Плутанина тут дає клієнту хибну суму."""

    class _Legacy:
        FEE_DEC_RATE = 0.006
        PLATA_STAVKA = 0.003
        PLATA_CAP_PDV = 11_980.80

        @staticmethod
        def money(x):
            return f"{int(round(float(x))):,}".replace(",", " ")

        @staticmethod
        def fee_amcu(amount, stage):
            return int(min(170_000, max(3_000, amount * 0.006))), "0,6%"

        @staticmethod
        def plata_za_podannya(v):
            povna = min(11_980.80, float(v) * 0.003 * 1.2)
            return round(povna, 2), round(povna / 3, 2)

        @staticmethod
        def ua_date(d):
            return f"{d.day}.{d.month}.{d.year}"

    def test_amcu_fee_is_taken_from_the_lot_value(self):
        f = E._fee_rows(4_045_480.0, 2_251_738.0, self._Legacy)
        self.assertIn("24 272", f["плата"])
        self.assertIn("ЛОТУ", f["плата"])
        self.assertIn("звірити чинність", f["плата"],
                      "ставку треба перевіряти, а не подавати як встановлену")

    def test_submission_fee_is_taken_from_the_bid_price(self):
        f = E._fee_rows(4_045_480.0, 2_251_738.0, self._Legacy)
        self.assertIn("8 106", f["втрачено"].replace("\xa0", " "),
                      "плата за подання рахується від ціни ПРОПОЗИЦІЇ")
        self.assertNotIn("4 045 480", f["втрачено"])

    def test_missing_bid_price_is_not_replaced_by_the_lot_value(self):
        """Правило 5: відсутнє число не підмінюється сусіднім."""
        f = E._fee_rows(4_045_480.0, None, self._Legacy)
        self.assertIn(E.NOT_ESTABLISHED, f["втрачено"])
        self.assertNotIn("11 980", f["втрачено"])
        self.assertIn("24 272", f["плата"], "плата до АМКУ рахується далі")

    def test_missing_lot_value_blocks_only_the_amcu_fee(self):
        f = E._fee_rows(None, 2_251_738.0, self._Legacy)
        self.assertIn(E.NOT_ESTABLISHED, f["плата"])
        self.assertIn("8 106", f["втрачено"].replace("\xa0", " "))

    def test_deadline_says_not_established_instead_of_staying_empty(self):
        self.assertIn(E.NOT_ESTABLISHED, E._deadline_row(None, self._Legacy))

    def test_expired_deadline_is_called_expired(self):
        row = E._deadline_row("2020-01-01T10:00:00+02:00", self._Legacy)
        self.assertIn("МИНУВ", row, "прострочений строк не можна показувати "
                                    "як залишок часу")


class TestCardShowsWhyOneLetter(Base):
    """Картка має пояснювати, чому на дві закупівлі пішов один лист."""

    def setUp(self):
        super().setUp()
        eng = self.engine([], mode="TEST")
        self.run_cycle(eng, TestNoDuplicateForTwoRejections.TWO)
        eng.send()
        eng.close()
        self.root = os.path.join(self.tmp, "drive")
        self.stats = E.build_case_cards(self.db_path, self.root, legacy=None)

    def _text(self, ua):
        import docx
        d = docx.Document(os.path.join(self.root, "spravy", ua,
                                       "00_KARTKA_12345678.docx"))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)

    def test_both_tender_ids_are_listed_in_the_card(self):
        text = self._text("UA-2026-08-07-008904-a")
        self.assertIn("UA-2026-08-07-009602-a", text,
                      "картка має показувати друге відхилення того ж учасника")
        self.assertIn("лист надіслано ОДИН", text)

    def test_the_card_without_a_letter_says_so(self):
        """Друга картка є, листа в ній немає — і це написано словами."""
        a = self._text("UA-2026-08-07-008904-a")
        b = self._text("UA-2026-08-07-009602-a")
        with_letter = [t for t in (a, b) if "Тема:" in t]
        without = [t for t in (a, b) if "Тема:" not in t]
        self.assertEqual(len(with_letter), 1, "лист має бути рівно в одній картці")
        self.assertIn("перший дотик уже був", without[0])

    def test_the_sent_letter_text_is_in_the_card(self):
        a = self._text("UA-2026-08-07-008904-a")
        b = self._text("UA-2026-08-07-009602-a")
        text = a if "Тема:" in a else b
        self.assertIn("ТЕКСТ НАДІСЛАНОГО ЛИСТА", text)
        self.assertIn("TenderWin", text)


# ================================================ словник Clarity (§15)
PROTO_4 = (
    "Тендерна пропозиція учасника ТОВ «Ірбудтранс» не відповідає вимогам "
    "п. 1 Розділу III тендерної документації, оскільки надана довідка про "
    "наявність працівників не містить інформації про інженерно-технічний "
    "персонал, чим порушено вимоги ч. 1 ст. 31 Закону України «Про публічні "
    "закупівлі» та п. 44 Особливостей, затверджених постановою КМУ від "
    "12.10.2022 № 1178. "
    "Крім того, учасник протягом 24 годин не усунув невідповідності, а саме "
    "не надав акти виконаних робіт за аналогічним договором, що вимагалось "
    "Додатком 3. "
    "Також строк дії банківської гарантії менший ніж вимагав п. 7 Додатку 4. "
    "Локальний кошторис не містить позицію 12 відомості обсягів робіт.")


class TestClaimAnalysis(Base):
    """§11 — розбір претензій. Головна вимога: жодна підстава не губиться."""

    def setUp(self):
        super().setUp()
        import lead_machine_v1 as L
        self.claims = E.analyze_protocol(PROTO_4, cpv="45000000-7", legacy=L)

    def test_every_ground_survives_the_splitter(self):
        """Правило 32: висновок неможливий, доки не розібрано ВСІ підстави."""
        self.assertEqual(len(self.claims), 4)
        codes = {c for cl in self.claims for c, _, _ in cl.codes}
        for expected in ("R01", "R04", "R03", "R06"):
            self.assertIn(expected, codes, f"підстава {expected} загубилась")

    def test_multi_label_not_first_match(self):
        """§11.1 — одна претензія може мати кілька тем."""
        multi = [c for c in self.claims if len(c.codes) > 1]
        self.assertTrue(multi, "жодна претензія не отримала кількох тем")

    def test_general_codes_do_not_swallow_precise_ones(self):
        """§15 п. 4 — R10/R99 не поглинають R06, R09, R01, R03."""
        for claim in self.claims:
            top = claim.codes[0][0]
            others = {c for c, _, _ in claim.codes}
            if others & {"R01", "R03", "R04", "R06", "R09"}:
                self.assertNotIn(top, ("R10", "R99"),
                                 f"{claim.claim_id}: загальний код поглинув точний")

    def test_citations_are_extracted_with_their_kind(self):
        c1 = self.claims[0]
        got = {(x["вид"], x["значення"]) for x in c1.citations}
        for expected in (("стаття", "31"), ("частина", "1"), ("розділ", "III"),
                         ("Особливості", "44"), ("ПКМУ", "1178")):
            self.assertIn(expected, got, f"не знайдено {expected}")

    def test_pp_does_not_double_count_as_p(self):
        c = E.analyze_claim("Кошторис не відповідає пп. 2.3 Додатку 4.", "C01")
        kinds = [(x["вид"], x["значення"]) for x in c.citations]
        self.assertIn(("підпункт", "2.3"), kinds)
        self.assertNotIn(("пункт", "2.3"), kinds, "«п.» спрацював усередині «пп.»")

    def test_missing_citation_is_a_finding_not_a_blank(self):
        """Відсутність посилання на пункт ТД — встановлений факт (X01)."""
        c = E.analyze_claim("Учасник не надав довідку про працівників.", "C01")
        self.assertEqual(c.citations, [])
        self.assertIn("X01_MOTIVATION", c.cross)

    def test_bare_norm_quote_does_not_become_a_24h_ground(self):
        """§11.2 п. 1 — цитата норми про 24 години не є претензією."""
        c = E.analyze_claim(
            "Замовник розміщує у système повідомлення з вимогою про усунення "
            "невідповідностей протягом 24 годин відповідно до Особливостей.",
            "C01")
        self.assertNotIn("R04", {x for x, _, _ in c.codes})

    def test_participant_name_never_leaks_into_the_query(self):
        """§7.2 — назви сторін і ЄДРПОУ в запиті шукали б лише цю справу."""
        for claim in self.claims:
            low = claim.exact_query.lower()
            for bad in ("ірбудтранс", "єдрпоу", "тов", "37406974"):
                self.assertNotIn(bad, low, f"{claim.claim_id}: «{bad}» у запиті")

    def test_query_length_is_within_the_rule(self):
        """§7.2 — 4–9 змістовних слів."""
        for claim in self.claims:
            n = len(claim.exact_query.split())
            self.assertLessEqual(n, 9, f"{claim.claim_id}: {n} слів")

    def test_weak_query_is_marked_not_hidden(self):
        c = E.analyze_claim("Пропозиція не відповідає документації.", "C01")
        self.assertEqual(c.exact_confidence, "LOW")
        self.assertTrue(any("LOW_CONFIDENCE" in n for n in c.notes))


class TestSearchOrder(Base):
    """§7.5, §15 — пошуковий наряд."""

    def setUp(self):
        super().setUp()
        import lead_machine_v1 as L
        self.claims = E.analyze_protocol(PROTO_4, cpv="45000000-7", legacy=L)
        self.rows = E.build_search_order(
            self.claims, cpv="45000000-7", title="Поточний ремонт",
            buyer_edrpou="01992682")

    def test_every_claim_gets_the_mandatory_pair(self):
        """§15 п. 5 — і дослівний запит, і КОНТРПОШУК, на кожну претензію."""
        for claim in self.claims:
            types = {r.query_type for r in self.rows
                     if r.claim_id == claim.claim_id}
            self.assertIn("exact_fact", types, claim.claim_id)
            self.assertIn("counter", types,
                          f"{claim.claim_id}: немає контрпошуку")

    def test_both_polarities_are_present(self):
        pol = {r.polarity for r in self.rows}
        self.assertIn("participant", pol)
        self.assertIn("buyer", pol)

    def test_limits_are_respected(self):
        """§7.5 — 8 на претензію, 20 на закупівлю."""
        self.assertLessEqual(len(self.rows), E.MAX_QUERIES_PER_TENDER)
        for claim in self.claims:
            n = len([r for r in self.rows if r.claim_id == claim.claim_id])
            self.assertLessEqual(n, E.MAX_QUERIES_PER_CLAIM)

    def test_budget_does_not_starve_the_last_ground(self):
        """Ліміт закупівлі не має з'їдати останні підстави цілком."""
        served = {r.claim_id for r in self.rows}
        self.assertEqual(served, {c.claim_id for c in self.claims},
                         "якась підстава лишилась без жодного запиту")

    def test_no_duplicate_queries(self):
        texts = [r.query_text.lower() for r in self.rows
                 if r.query_type != "buyer"]
        self.assertEqual(len(texts), len(set(texts)))

    def test_sector_terms_only_when_the_text_confirms_them(self):
        """§10 — не чіпляти кошторисну лексику до спору про гарантію."""
        for r in self.rows:
            if r.query_type != "sector":
                continue
            claim = next(c for c in self.claims if c.claim_id == r.claim_id)
            self.assertTrue(claim.sector_confirmed,
                            f"{r.claim_id}: галузевий запит без підстав у тексті")

    def test_filters_are_printed_not_pressed(self):
        f = self.rows[0].filters
        self.assertEqual(f["класифікація"], "роботи")
        self.assertIn("Prozorro", f["джерело"])

    def test_buyer_pass_is_separate(self):
        """§6.5 — пошук за замовником є окремим проходом, а не основним."""
        buyer_rows = [r for r in self.rows if r.query_type == "buyer"]
        self.assertTrue(buyer_rows)
        self.assertEqual(buyer_rows[0].filters["ЄДРПОУ замовника"], "01992682")

    def test_nothing_reaches_clarity_over_the_network(self):
        """§2 — скрипт до Clarity не звертається. Перевіряємо код модуля."""
        import inspect
        src = "".join(inspect.getsource(fn) for fn in
                      (E.build_search_order, E.analyze_claim, E.analyze_protocol,
                       E.keywords_line, E.search_order_markdown))
        for bad in ("requests.", "urlopen", "http://", "https://", "session"):
            self.assertNotIn(bad, src, f"у генераторі знайдено «{bad}»")

    def test_no_forbidden_conclusions_anywhere(self):
        """§14 — скрипт не має права оцінювати перспективу."""
        text = " ".join([r.query_text + " " + r.why for r in self.rows])
        text += " " + E.keywords_line(self.claims, self.rows)
        text += " " + E.search_order_markdown(self.rows, self.claims, "UA-X")
        low = text.lower()
        for bad in ("шанс", "перспектива висока", "виграш", "гарантує",
                    "прецедент", "завжди стає"):
            self.assertNotIn(bad, low, f"заборонене формулювання: {bad}")

    def test_keywords_line_meets_the_minimum(self):
        kw = E.keywords_line(self.claims, self.rows).splitlines()
        self.assertGreaterEqual(len(kw), 22, "у картці має бути не менше 22")
        self.assertLessEqual(len(kw), 26)


class TestCardDetail(Base):
    """Картка: повні підстави, посилання на протокол, наряд."""

    def setUp(self):
        super().setUp()
        import lead_machine_v1 as L
        self.L = L
        eng = self.engine([], mode="TEST", legacy=L)
        self.run_cycle(eng, [facts(reason=PROTO_4)])
        eng.send()
        eng.db.conn.execute(
            "UPDATE rejections SET cpv='45000000-7', lot_amount=4045480.0,"
            " bid_amount=2251738.0, winner_amount=2609900.0,"
            " winner_name='ТОВ «Переможець»', notice_24h='YES',"
            " protocol_url='https://public.docs.openprocurement.org/get/abc',"
            " protocol_title='Протокол.pdf'")
        eng.db.conn.commit()
        eng.close()
        self.root = os.path.join(self.tmp, "drive")
        self.stats = E.build_case_cards(self.db_path, self.root, legacy=L)
        self.folder = os.path.join(self.root, "spravy",
                                   "UA-2026-08-28-000001-a")

    def _doc(self):
        import docx
        return docx.Document(os.path.join(self.folder, "00_KARTKA_12345678.docx"))

    def _text(self):
        d = self._doc()
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)

    def test_every_ground_is_written_out_in_full(self):
        text = self._text()
        for cid in ("C01", "C02", "C03", "C04"):
            self.assertIn(cid, text, f"підстави {cid} немає в картці")
        self.assertIn("довідка про наявність працівників", text)
        self.assertIn("банківської гарантії", text)
        self.assertIn("Локальний кошторис", text)

    def test_protocol_link_is_clickable(self):
        d = self._doc()
        targets = [r.target_ref for r in d.part.rels.values() if r.is_external]
        self.assertTrue(any("openprocurement" in t for t in targets),
                        "немає гіперпосилання на протокол")

    def test_short_reason_row_is_replaced_by_the_detail(self):
        text = self._text()
        self.assertNotIn("Підстава — коротко в листі", text)
        self.assertIn("ЩО ПЕРЕВІРИТИ ПЕРЕД РОЗМОВОЮ", text)
        self.assertIn("норми в цій підставі", text)

    def test_clarity_keywords_are_gone_from_the_card(self):
        # Рішення власника продукту 2026-09-02: пошукові слова Clarity
        # з картки прибрані. Генератор лишається в коді за прапорцем.
        self.assertNotIn("КЛЮЧОВІ СЛОВА", self._text().upper())
        self.assertFalse(os.path.isfile(
            os.path.join(self.folder, "00_ZAPYTY_CLARITY_12345678.md")))
        self.assertFalse(E.CLARITY_KEYWORDS_IN_CARD)

    def test_price_gap_to_the_winner_is_shown(self):
        self.assertIn("ДЕШЕВШЕ за переможця", self._text())

    def test_unknown_price_is_not_invented(self):
        eng = self.engine([], mode="TEST", legacy=self.L,
                          db_path=os.path.join(self.tmp, "b.db"))
        self.run_cycle(eng, [facts(reason=PROTO_4)])
        eng.close()
        root2 = os.path.join(self.tmp, "d2")
        E.build_case_cards(os.path.join(self.tmp, "b.db"), root2, legacy=self.L)
        import docx
        p = os.path.join(root2, "spravy", "UA-2026-08-28-000001-a",
                         "00_KARTKA_12345678.docx")
        d = docx.Document(p)
        rows = [" | ".join(c.text for c in r.cells)
                for t in d.tables for r in t.rows]
        gap = [r for r in rows if r.startswith("Різниця з переможцем")]
        self.assertTrue(gap)
        self.assertIn(E.NOT_ESTABLISHED, gap[0])
