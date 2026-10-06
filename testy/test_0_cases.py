# -*- coding: utf-8 -*-
"""
Тести тек справ «0 Cases» (Lead & Mail 1.4.0, Service Cards 1.1.0).

Рішення власника 06.10.2026: службова картка і протокол (рішення)
замовника лягають у «G:\\Мій диск\\0 Cases» (у Colab —
/content/drive/MyDrive/0 Cases) у підтеку «<ID закупівлі> <ЄДРПОУ/ІПН>».

Мережі немає: документи віддає підмінений завантажувач.

Запуск:  python3 test_0_cases.py   (поруч із модулями і test_lead_mail.py)
"""
from __future__ import annotations

import io
import json
import os
import shutil
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import tenderwin_lead_mail as M              # noqa: E402
import tenderwin_service_cards as SC         # noqa: E402
import test_lead_mail as T                   # noqa: E402  фікстури Prozorro

UA = "UA-2026-10-02-004417-a"
CODE = "41234567"


def protocol_docx(text: str = "Пропозицію відхилено: учасник не надав довідку.") -> bytes:
    import docx                                                  # noqa: PLC0415
    document = docx.Document()
    document.add_paragraph("ПРОТОКОЛ № 45")
    document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class FakeFetcher:
    """Завантажувач без мережі: рахує, скільки разів його кликали."""

    def __init__(self, data: bytes = b""):
        self.data = data or protocol_docx()
        self.calls = []

    def __call__(self, url: str) -> SC.Fetched:
        self.calls.append(url)
        return SC.Fetched(data=self.data, http_status=200)


def row(object_id="award-0001-aaaa", ua_id=UA, code=CODE, event_time="2026-10-02T11:20:00+03:00",
        docs=None, **extra) -> dict:
    base = {
        "event_id": f"ev-{object_id}", "ua_id": ua_id, "tender_id": f"tid-{ua_id}",
        "title": "Капітальний ремонт покрівлі", "buyer_name": "Управління освіти",
        "buyer_edrpou": "00000001", "company_name": "ТОВ «БУДРЕМСЕРВІС ПЛЮС»",
        "edrpou_norm": code, "stage": "awards", "object_id": object_id,
        "event_time": event_time, "object_status": "unsuccessful",
        "reason_raw": "Відхилення\nУчасник не надав довідку.",
        "docs_decision": json.dumps(docs if docs is not None else [
            {"id": f"doc-{object_id}", "title": "Протокол № 45.docx",
             "url": f"https://example/{object_id}/protocol",
             "published_iso": "2026-10-02T11:19:00+03:00"}], ensure_ascii=False),
        "state": "NEW", "complaint_end": "2026-10-07T00:00:00+03:00",
    }
    base.update(extra)
    return base


class CasesTmp(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="cases_")
        self.cases = os.path.join(self.tmp, "MyDrive", "0 Cases")
        self.index = os.path.join(self.tmp, "MyDrive", "TenderWin", "Службові картки")

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def build(self, rows, fetcher=None, label="R1"):
        return SC.build_service_cards(rows, root=self.cases, index_root=self.index,
                                      run_label=label, doc_fetcher=fetcher or FakeFetcher(),
                                      verbose=False, workers=4)


class TestFolderName(unittest.TestCase):
    def test_id_and_code(self):
        self.assertEqual(SC.case_folder_name({"ua_id": UA, "edrpou_norm": CODE}),
                         f"{UA} {CODE}")

    def test_individual_entrepreneur_ipn(self):
        self.assertEqual(SC.case_folder_name({"ua_id": UA, "edrpou_norm": "1234567890"}),
                         f"{UA} 1234567890")

    def test_windows_forbidden_characters_are_replaced(self):
        name = SC.case_folder_name({"ua_id": 'UA:1/2', "edrpou_norm": 'a*b?"'})
        for bad in '<>:"/\\|?*':
            self.assertNotIn(bad, name)

    def test_missing_values_are_named_not_hidden(self):
        self.assertEqual(SC.case_folder_name({}), "bez-UA bez-kodu")


class TestCaseFolders(CasesTmp):
    def test_card_and_protocol_land_in_the_case_folder(self):
        stats = self.build([row()])
        folder = os.path.join(self.cases, f"{UA} {CODE}")
        files = sorted(os.listdir(folder))
        self.assertIn(f"Службова_картка_{UA}.docx", files)
        self.assertIn("01_Протокол_№_45.docx", files)
        self.assertIn("manifest.json", files)
        self.assertEqual(stats["карток"]["CREATED"], 1)
        self.assertEqual(stats["тека"], self.cases)

    def test_cases_folder_holds_only_cases(self):
        """У «0 Cases» — лише теки справ: без тек дат і без переліків."""
        stats = self.build([row(), row("award-0002-bbbb", ua_id="UA-2026-10-02-000111-b",
                                       code="1234567890")])
        self.assertEqual(sorted(os.listdir(self.cases)),
                         ["UA-2026-10-02-000111-b 1234567890", f"{UA} {CODE}"])
        self.assertEqual(len(stats["перелік"]), 1)
        self.assertTrue(stats["перелік"][0].startswith(os.path.join(self.index, "2026-10-02")))
        self.assertTrue(os.path.exists(stats["перелік"][0]))

    def test_protocol_text_is_in_the_card(self):
        import docx                                              # noqa: PLC0415
        self.build([row()])
        card = docx.Document(os.path.join(self.cases, f"{UA} {CODE}",
                                          f"Службова_картка_{UA}.docx"))
        text = "\n".join(p.text for p in card.paragraphs)
        text += "\n".join(c.text for t in card.tables for r in t.rows for c in r.cells)
        self.assertIn("учасник не надав довідку", text)
        self.assertIn(os.path.join("0 Cases", f"{UA} {CODE}"), text)

    def test_index_points_to_the_case_folder(self):
        import openpyxl                                          # noqa: PLC0415
        stats = self.build([row()])
        sheet = openpyxl.load_workbook(stats["перелік"][0]).active
        header = [c.value for c in sheet[1]]
        column = header.index("Тека справи")
        self.assertEqual(sheet.cell(row=2, column=column + 1).value,
                         os.path.join("0 Cases", f"{UA} {CODE}"))

    def test_rerun_takes_the_same_folder_and_does_not_download_again(self):
        self.build([row()])
        again = FakeFetcher()
        stats = self.build([row()], fetcher=again, label="R2")
        self.assertEqual(again.calls, [])
        self.assertEqual(os.listdir(self.cases), [f"{UA} {CODE}"])
        self.assertEqual(stats["карток"]["UPDATED"], 1)

    def test_second_decision_of_the_same_company_gets_its_own_folder(self):
        """Два лоти або повторне рішення: друга справа не затирає першу."""
        first = row("award-0001-aaaa", event_time="2026-10-02T11:20:00+03:00")
        second = row("award-0002-bbbb", event_time="2026-10-02T15:00:00+03:00")
        self.build([second, first])
        self.assertEqual(sorted(os.listdir(self.cases)),
                         [f"{UA} {CODE}", f"{UA} {CODE} (2)"])
        owner = SC.folder_owner(os.path.join(self.cases, f"{UA} {CODE} (2)"))
        self.assertEqual(owner, "award-0002-bbbb")
        # Повтор лише другого рішення знаходить його теку, а не першу.
        again = FakeFetcher()
        self.build([second], fetcher=again, label="R2")
        self.assertEqual(again.calls, [])
        self.assertEqual(len(os.listdir(self.cases)), 2)
        self.assertEqual(SC.folder_owner(os.path.join(self.cases, f"{UA} {CODE}")),
                         "award-0001-aaaa")

    def test_files_of_a_person_in_the_case_folder_are_kept(self):
        folder = os.path.join(self.cases, f"{UA} {CODE}")
        os.makedirs(folder)
        with open(os.path.join(folder, "мій_аналіз.txt"), "w", encoding="utf-8") as fh:
            fh.write("нотатки")
        self.build([row()])
        with open(os.path.join(folder, "мій_аналіз.txt"), encoding="utf-8") as fh:
            self.assertEqual(fh.read(), "нотатки")
        self.assertIn(f"Службова_картка_{UA}.docx", os.listdir(folder))

    def test_assignment_is_by_decision_not_by_order(self):
        rows = [row("award-0002-bbbb", event_time="2026-10-02T15:00:00+03:00"),
                row("award-0001-aaaa", event_time="2026-10-02T11:20:00+03:00")]
        self.assertEqual(SC.assign_case_folders(rows, self.cases),
                         [f"{UA} {CODE} (2)", f"{UA} {CODE}"])


class TestCasesRoot(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="root_")
        self.saved = {k: getattr(M, k) for k in ("CASES_DIR", "WORK_DIR", "DB_PATH",
                                                 "OUT_DIR", "LEGACY_DB_PATH",
                                                 "CARDS_DIR", "MY_DRIVE_ROOTS")}
        self.token_dir = M.E.TOKEN_DIR

    def tearDown(self):
        for key, value in self.saved.items():
            setattr(M, key, value)
        M.E.TOKEN_DIR = self.token_dir
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_colab_my_drive(self):
        dysk = os.path.join(self.tmp, "content", "drive", "MyDrive", "Tenderwin")
        M._paths(dysk)
        self.assertEqual(M.cases_root(),
                         os.path.join(self.tmp, "content", "drive", "MyDrive", "0 Cases"))

    def test_my_drive_in_ukrainian(self):
        dysk = os.path.join(self.tmp, "Мій диск", "Робота", "TenderWin")
        M._paths(dysk)
        self.assertEqual(M.cases_root(), os.path.join(self.tmp, "Мій диск", "0 Cases"))

    def test_explicit_setting_wins(self):
        M._paths(os.path.join(self.tmp, "MyDrive", "Tenderwin"))
        M.CASES_DIR = os.path.join(self.tmp, "інша", "0 Cases")
        self.assertEqual(M.cases_root(), M.CASES_DIR)

    def test_known_mount_point_when_work_folder_is_elsewhere(self):
        drive = os.path.join(self.tmp, "G", "My Drive")
        os.makedirs(drive)
        M.MY_DRIVE_ROOTS = (os.path.join(self.tmp, "немає"), drive)
        M._paths(os.path.join(self.tmp, "local", "TenderWin"))
        self.assertEqual(M.cases_root(), os.path.join(drive, "0 Cases"))

    def test_without_any_drive_the_cases_stay_inside_the_work_folder(self):
        M.MY_DRIVE_ROOTS = (os.path.join(self.tmp, "немає"),)
        dysk = os.path.join(self.tmp, "local", "TenderWin")
        M._paths(dysk)
        self.assertEqual(M.cases_root(), os.path.join(dysk, "0 Cases"))


class TestLeadMailWritesToCases(T.Base):
    """Повний шлях Lead & Mail: подія з бази → тека справи в «0 Cases»."""

    def setUp(self):
        super().setUp()
        self.saved_cases = (M.CASES_DIR, M.CARDS_DIR)
        M.CASES_DIR = os.path.join(self.tmp, "MyDrive", "0 Cases")
        M.CARDS_DIR = os.path.join(self.tmp, "TenderWin", "Службові картки")

    def tearDown(self):
        M.CASES_DIR, M.CARDS_DIR = self.saved_cases
        super().tearDown()

    def test_event_of_the_run_gets_a_case_folder_with_the_protocol(self):
        docs = [{"id": "p1", "title": "Протокол № 45.docx", "url": "https://example/p1",
                 "datePublished": T.iso(22, 11), "author": "tender_owner"}]
        t = T.tender(bids=[T.bid("b1")],
                     awards=[T.award("a1", "b1", date=T.iso(22, 11), cp_start=T.iso(22, 11),
                                     cp_end=T.iso(27, 0), docs=docs)])
        eng, _ = self.run_scan([t])
        fetcher = FakeFetcher()
        stats = eng.service_cards(run_id=eng.run_id, fresh=False, doc_fetcher=fetcher)
        eng.close()
        folder = os.path.join(M.CASES_DIR, "UA-2026-09-01-000001-a 12345678")
        self.assertTrue(os.path.isdir(folder), os.listdir(M.CASES_DIR))
        files = os.listdir(folder)
        self.assertIn("Службова_картка_UA-2026-09-01-000001-a.docx", files)
        self.assertIn("01_Протокол_№_45.docx", files)
        self.assertEqual(fetcher.calls, ["https://example/p1"])
        self.assertTrue(all(p.startswith(M.CARDS_DIR) for p in stats["перелік"]))
        self.assertEqual(os.listdir(M.CASES_DIR), ["UA-2026-09-01-000001-a 12345678"])


if __name__ == "__main__":
    unittest.main(verbosity=1)
