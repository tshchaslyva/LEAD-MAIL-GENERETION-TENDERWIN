# -*- coding: utf-8 -*-
"""
Тести редакції листа і повного набору документів (2026-09-02).
Мережа не потрібна: завантажувач і транспорт підмінюються.
"""
from __future__ import annotations

import io
import json
import os
import shutil
import sqlite3
import sys
import tempfile
import types
import unittest

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import tenderwin_lead_engine as E   # noqa: E402

E.ALLOW_COLAB_SECRETS = False

import lead_machine_v1 as L         # noqa: E402
import test_lead_engine as T        # noqa: E402


class TestSubjectGenitive(unittest.TestCase):
    """«на закупівлю поточного ремонту», а не «Поточний ремонт»."""

    KNOWN = [
        ("Поточний ремонт приміщень будівлі", "поточного ремонту приміщень"),
        ("Капітальний ремонт даху", "капітального ремонту даху"),
        ("Реконструкція будівлі садка", "реконструкції будівлі"),
        ("Нове будівництво мережі", "нового будівництва мережі"),
        ("Послуги з технічного нагляду", "послуг з технічного нагляду"),
        ("Будівельні роботи з облаштування", "будівельних робіт"),
        ("Термомодернізація будівлі ліцею", "термомодернізації будівлі"),
        ("Заміна вікон", "заміни вікон"),
    ]

    def test_known_heads_are_declined(self):
        for title, expected in self.KNOWN:
            with self.subTest(title=title):
                text, conf = E.subject_genitive(title)
                self.assertEqual(conf, "HIGH")
                self.assertIn(expected, text)

    def test_dk_code_is_stripped(self):
        text, _ = E.subject_genitive("ДК 021:2015: 45453000-7 — Капітальний ремонт")
        self.assertNotIn("021", text)
        self.assertTrue(text.startswith("капітального ремонту"))

    def test_coordinated_head_is_declined_too(self):
        text, _ = E.subject_genitive("Капітальний ремонт і реставрація")
        self.assertIn("реставрації", text, "другий іменник лишився в називному")

    def test_unknown_head_is_quoted_not_invented(self):
        """Правило 5: невідому форму не вигадуємо."""
        text, conf = E.subject_genitive("Ремонтно-відновлювальні заходи")
        self.assertEqual(conf, "LOW")
        self.assertTrue(text.startswith("«") and text.endswith("»"))

    def test_empty_title_is_empty(self):
        self.assertEqual(E.subject_genitive("")[1], "NONE")


class TestBuyerAgreement(unittest.TestCase):
    """«яку проводило Управління», а не «яку проводив Управління»."""

    CASES = [
        ("Управління освіти", "проводило"),
        ("Департамент будівництва", "проводив"),
        ("Служба автомобільних доріг", "проводила"),
        ("КНП «Тячівська лікарня»", "проводило"),
        ("Одеська обласна рада", "проводила"),
        ("Виконавчий комітет міської ради", "проводив"),
    ]

    def test_verb_agrees_with_the_buyer(self):
        for buyer, verb in self.CASES:
            with self.subTest(buyer=buyer):
                self.assertEqual(E.buyer_verb(buyer)[0], verb)

    def test_full_legal_form_is_abbreviated(self):
        self.assertEqual(
            E.short_org_name("ТОВАРИСТВО З ОБМЕЖЕНОЮ ВІДПОВІДАЛЬНІСТЮ «АЛЬФА»"),
            "ТОВ «АЛЬФА»")
        self.assertEqual(
            E.short_org_name("КОМУНАЛЬНЕ НЕКОМЕРЦІЙНЕ ПІДПРИЄМСТВО «РЛ»"),
            "КНП «РЛ»")

    def test_short_name_is_left_alone(self):
        self.assertEqual(E.short_org_name("ТОВ «АЛЬФА»"), "ТОВ «АЛЬФА»")


class LetterBase(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.db = os.path.join(self.tmp, "leads.db")

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def letter(self, **kw):
        cfg = E.Config(db_path=self.db, out_dir=os.path.join(self.tmp, "o"),
                       mode="TEST", min_delay=0, max_delay=0, verbose=False)
        eng = E.LeadEngine(cfg, E.FixtureSource([]),
                           E.DryRunTransport(os.path.join(self.tmp, "e")), L)
        eng.repo.mark_history_migrated("тест")
        eng.ingest([T.facts(**kw)])
        eng.build_queue()
        row = eng.db.one("SELECT subject_rendered, body_rendered, template_version"
                         " FROM outreach_events")
        eng.close()
        return dict(row) if row else {}


class TestLetterV2(LetterBase):
    """Канонічна редакція First Touch v3 від 2026-09-03."""

    def test_subject_is_only_the_procurement_id(self):
        """§2: тема — «Щодо закупівлі UA-...», без замовника і без реклами."""
        row = self.letter()
        self.assertEqual(row["subject_rendered"],
                         "Щодо закупівлі UA-2026-08-28-000001-a")
        for banned in ("Терміново", "АМКУ", "Оскарження", "Виграємо",
                       "Комерційна пропозиція", "Важлива інформація"):
            self.assertNotIn(banned, row["subject_rendered"])

    def test_greeting_uses_the_given_name_only(self):
        """§8: «Добрий день, Іване!», а не повне ПІБ."""
        body = self.letter(name="Іван Петрович Коваленко")["body_rendered"]
        self.assertTrue(body.startswith("Добрий день, Іване!"), body[:40])
        self.assertNotIn("Коваленко", body)

    def test_greeting_without_a_name_is_plain(self):
        body = self.letter(name="", company="ТОВ «Альфа»")["body_rendered"]
        self.assertTrue(body.startswith("Добрий день!"))
        self.assertNotIn("Добрий день,", body.splitlines()[0])

    def test_company_name_in_the_field_is_not_declined(self):
        body = self.letter(name="Тов Ірбудтранс",
                           company="ТЗОВ «ІРБУДТРАНС»")["body_rendered"]
        self.assertTrue(body.startswith("Добрий день!"))
        self.assertNotIn("Ірбудтрансе", body)

    def test_the_body_follows_the_canonical_wording(self):
        body = self.letter()["body_rendered"]
        for phrase in ("Побачив, що пропозицію",
                       "не поспішати зі скаргою",
                       "чи справді у замовника були достатні підстави",
                       "надішлю свій погляд на ситуацію",
                       "З повагою,", "Віталій Щасливий",
                       "радник з публічних закупівель у будівництві",
                       "TenderWin"):
            self.assertIn(phrase, body, f"немає фрази «{phrase}»")
        self.assertIn(E.SENDER_PHONE, body)

    def test_no_sales_call_to_action(self):
        """§20: у першому листі не має бути жодного продажного заклику."""
        body = self.letter()["body_rendered"].lower()
        for cta in ("замовити консультац", "зателефонуйте", "купити аналіз",
                    "подати скаргу", "забронювати", "ціна", "вартість послуг",
                    "знижк", "акці"):
            self.assertNotIn(cta, body, f"продажний заклик: {cta}")

    def test_no_placeholders_survive(self):
        """§51: якщо мітка лишилась — лист не має піти."""
        body = self.letter()["body_rendered"]
        self.assertEqual(E._leftover_placeholders(body), [])

    def test_same_day_promise_inside_the_working_window(self):
        import datetime as dt
        morning = dt.datetime(2026, 9, 3, 10, 0, tzinfo=E.KYIV)
        self.assertIn("вже сьогодні", E.followup_sentence(morning))

    def test_no_same_day_promise_after_the_cutoff(self):
        """§4: обіцянку «сьогодні» не можна давати наосліп."""
        import datetime as dt
        evening = dt.datetime(2026, 9, 3, 19, 0, tzinfo=E.KYIV)
        text = E.followup_sentence(evening)
        self.assertNotIn("сьогодні", text)
        self.assertIn("після перевірки", text)

    def test_no_same_day_promise_on_a_weekend(self):
        import datetime as dt
        saturday = dt.datetime(2026, 9, 5, 10, 0, tzinfo=E.KYIV)
        self.assertNotIn("сьогодні", E.followup_sentence(saturday))

    def test_the_gate_can_be_switched_off_entirely(self):
        E.SAME_DAY_FOLLOWUP_ENABLED = False
        try:
            import datetime as dt
            morning = dt.datetime(2026, 9, 3, 10, 0, tzinfo=E.KYIV)
            self.assertNotIn("сьогодні", E.followup_sentence(morning))
        finally:
            E.SAME_DAY_FOLLOWUP_ENABLED = True

    def test_old_versions_stay_in_the_database(self):
        """§19: історичні версії шаблонів не переписуються."""
        row = self.letter()
        self.assertEqual(row["template_version"], "V3")
        con = sqlite3.connect(self.db)
        versions = {(r[0], r[1]) for r in con.execute(
            "SELECT template_id, template_version FROM templates")}
        con.close()
        self.assertIn(("TENDERWIN_FIRST_TOUCH_V3", "V3"), versions)
        self.assertIn(("TENDERWIN_FIRST_TOUCH_B", "V1"), versions)
        self.assertIn(("TENDERWIN_FIRST_TOUCH_B", "V2"), versions)

    def test_html_alternative_is_clean(self):
        """§13, §18: без банерів, кнопок, зображень і піксельного стеження."""
        html = E.letter_html(self.letter()["body_rendered"])
        self.assertIn("<p style=", html)
        self.assertIn('href="https://tenderwin.com.ua"', html)
        self.assertIn("Arial", html)
        for forbidden in ("<script", "<img", "<form", "<button", "<table",
                          "background-image", "googleapis", "px-tracking"):
            self.assertNotIn(forbidden, html.lower(), f"заборонено: {forbidden}")

    def test_html_escapes_company_names(self):
        html = E.letter_html('Добрий день!\n\nТОВ <b>"Х"</b> & Ко\n\nЗ повагою,\nВіталій')
        self.assertIn("&lt;b&gt;", html)
        self.assertIn("&amp;", html)

    def test_no_leftover_placeholders(self):
        body = self.letter()["body_rendered"]
        self.assertNotIn("{", body)
        self.assertNotIn("[", body.split("З повагою")[0])


def _xlsx_bytes():
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["№", "Найменування робіт", "Од.", "Кількість", "Ціна"])
    ws.append([1, "Розбирання покриття підлоги", "м2", 125.40, 87.65])
    ws.append([2, "Улаштування стяжки", "м2", 125.04, 214.30])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _docx_bytes():
    import docx
    d = docx.Document()
    d.add_paragraph("Довідка про наявність працівників")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


class DocsBase(unittest.TestCase):
    TD = [{"назва": "Тендерна документація.pdf", "url": "https://p/td1",
           "опубліковано": "2026-07-01"},
          {"назва": "Тендерна документація.pdf", "url": "https://p/td1b",
           "опубліковано": "2026-07-18"},
          {"назва": "Додаток 2.docx", "url": "https://p/td2"}]
    BID = [{"назва": "Довідка.docx", "url": "https://p/b1"},
           {"назва": "Кошторис.xlsx", "url": "https://p/b2"},
           {"назва": "Гарантія.pdf", "url": "https://p/b3"}]
    AW = [{"назва": "Протокол.pdf", "url": "https://p/proto"}]
    QA = [{"назва": "Звернення.pdf", "url": "https://p/q1"}]
    QA_TEXT = ("[ЗАПИТАННЯ] 2026-07-20 · статус: answered\n"
               "  питання: Чи потрібна довідка окремим файлом?\n"
               "  відповідь: Так, окремим файлом.")

    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.db = os.path.join(self.tmp, "leads.db")
        cfg = E.Config(db_path=self.db, out_dir=os.path.join(self.tmp, "o"),
                       mode="TEST", min_delay=0, max_delay=0, verbose=False)
        eng = E.LeadEngine(cfg, E.FixtureSource([]),
                           E.DryRunTransport(os.path.join(self.tmp, "e")), L)
        eng.repo.mark_history_migrated("тест")
        eng.ingest([T.facts(ua="UA-2026-08-07-008904-a")])
        eng.build_queue()
        eng.close()
        con = sqlite3.connect(self.db)
        con.execute(
            "UPDATE rejections SET td_documents=?, bid_documents=?,"
            " award_documents=?, qa_documents=?, qa_text=?",
            (json.dumps(self.TD, ensure_ascii=False),
             json.dumps(self.BID, ensure_ascii=False),
             json.dumps(self.AW, ensure_ascii=False),
             json.dumps(self.QA, ensure_ascii=False), self.QA_TEXT))
        con.commit()
        con.close()
        self.bytes = {"https://p/b1": _docx_bytes(),
                      "https://p/b2": _xlsx_bytes(),
                      "https://p/td2": _docx_bytes()}
        self.net = types.SimpleNamespace(
            make_session=lambda: object(),
            get_bytes=lambda s, u, tries=3: self.bytes.get(u, b"%PDF-1.4 x"))
        self.case = E.run_folder(self.tmp, "run-test-abc")
        # Документи живуть в АРХІВІ справи, а не в теці прогону: інакше
        # кожен новий прогін бачив порожній маніфест і качав усе заново.
        self.folder = E.case_archive(self.tmp, "UA-2026-08-07-008904-a")

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def download(self, **kw):
        return E.download_case_documents(self.db, self.tmp, legacy=self.net,
                                         case_dir=self.case, **kw)

    def files(self):
        return sorted(os.listdir(self.folder))


class TestAllDocuments(DocsBase):
    """«ВСІ документи пропозиції, вся ТД зі змінами, рішення, 24 год, запити»."""

    def test_every_group_lands_in_the_folder(self):
        self.download(make_pdf_copies=False)
        names = self.files()
        for prefix, what in (("01_rishennia", "рішення замовника"),
                             ("03_TD", "тендерна документація"),
                             ("04_propozycia", "пропозиція учасника"),
                             ("05_lystuvannia", "листування із замовником")):
            self.assertTrue(any(n.startswith(prefix) for n in names),
                            f"немає групи «{what}»: {names}")

    def test_both_editions_of_the_tender_documents_are_kept(self):
        self.download(make_pdf_copies=False)
        td = [n for n in self.files() if n.startswith("03_TD")
              and "Тендерна документація" in n]
        self.assertEqual(len(td), 2, "зміни до ТД мають зберігатись окремо")
        self.assertTrue(any("2026-07-18" in n for n in td),
                        "редакції не відрізняються датою")

    def test_questions_and_answers_are_saved_as_text(self):
        self.download(make_pdf_copies=False)
        path = [n for n in self.files() if n.endswith("zapytannya_ta_vidpovidi.txt")]
        self.assertTrue(path, "листування із замовником не збережено")
        with open(os.path.join(self.folder, path[0]), encoding="utf-8") as fh:
            body = fh.read()
        self.assertIn("Чи потрібна довідка окремим файлом", body)
        self.assertIn("Так, окремим файлом", body)

    def test_the_manifest_lists_everything(self):
        self.download(make_pdf_copies=False)
        with open(os.path.join(self.folder, E.MANIFEST_NAME),
                  encoding="utf-8") as fh:
            man = json.load(fh)
        self.assertEqual(len(man), 8, f"у маніфесті {len(man)} записів")


class TestPdfCopies(DocsBase):
    """PDF-копії для форматів, яких NotebookLM не читає."""

    def test_office_files_get_a_pdf_copy(self):
        """
        Excel NotebookLM НЕ читає — копія обовʼязкова.
        Word (.docx) з листопада 2025 NotebookLM читає сам — копія зайва і
        лише забирає одне з 50 місць для джерел. Звірено з офіційною
        сторінкою support.google.com/notebooklm/answer/16215270 (20.09.2026).
        """
        stats = self.download()
        names = self.files()
        self.assertTrue(
            any("Кошторис" in n and n.endswith(E.PDF_COPY_SUFFIX) for n in names),
            f"немає PDF-копії для Кошторис.xlsx: {names}")
        self.assertFalse(
            any("Довідка" in n and n.endswith(E.PDF_COPY_SUFFIX) for n in names),
            f"для .docx копія зайва — NotebookLM читає його сам: {names}")
        self.assertGreaterEqual(stats["pdf_копій"], 1)

    def test_pdf_originals_are_not_duplicated(self):
        self.download()
        copies = [n for n in self.files() if n.endswith(E.PDF_COPY_SUFFIX)]
        self.assertFalse(any("Гарантія" in n for n in copies),
                         "для PDF копія не потрібна")

    def test_numbers_survive_the_conversion(self):
        """Правило 19: кількісні показники — предмет спору, їх не можна втратити."""
        if E._soffice_bin() is None:
            self.skipTest("LibreOffice недоступний")
        self.download()
        pdf = [n for n in self.files()
               if "Кошторис" in n and n.endswith(E.PDF_COPY_SUFFIX)]
        self.assertTrue(pdf)
        import subprocess
        text = subprocess.run(
            ["pdftotext", "-layout", os.path.join(self.folder, pdf[0]), "-"],
            capture_output=True, text=True).stdout
        self.assertIn("125.4", text)
        self.assertIn("125.04", text)
        self.assertNotEqual("125.4" in text and "125.04" in text, False,
                            "125,40 і 125,04 мають лишатись різними")

    def test_long_cell_text_is_not_clipped(self):
        if E._soffice_bin() is None:
            self.skipTest("LibreOffice недоступний")
        self.download()
        pdf = [n for n in self.files()
               if "Кошторис" in n and n.endswith(E.PDF_COPY_SUFFIX)][0]
        import subprocess
        text = subprocess.run(
            ["pdftotext", "-layout", os.path.join(self.folder, pdf), "-"],
            capture_output=True, text=True).stdout
        self.assertIn("Розбирання покриття підлоги", text,
                      "текст обрізано по ширині колонки")

    def test_missing_converter_is_reported_not_hidden(self):
        real = E._soffice_bin
        E._soffice_bin = lambda: None
        try:
            # .xlsx NotebookLM не читає — отже, без LibreOffice копії не буде,
            # і це треба СКАЗАТИ, а не промовчати.
            path = os.path.join(self.tmp, "x.xlsx")
            with open(path, "wb") as fh:
                fh.write(_xlsx_bytes())
            pdf, status = E.convert_to_pdf(path)
            self.assertIsNone(pdf)
            self.assertEqual(status, "НЕМАЄ_LIBREOFFICE")
        finally:
            E._soffice_bin = real

    def test_unknown_format_says_so(self):
        path = os.path.join(self.tmp, "arkhiv.zip")
        with open(path, "wb") as fh:
            fh.write(b"PK\x03\x04")
        self.assertEqual(E.convert_to_pdf(path)[1], "ФОРМАТ_НЕ_КОНВЕРТУЄТЬСЯ")


class TestRunFolders(DocsBase):
    """«Спершу тека з датою, у ній — підпапка прогону»."""

    def test_path_is_date_then_run(self):
        import datetime as dt
        when = dt.datetime(2026, 9, 2, 14, 30, tzinfo=E.KYIV)
        path = E.run_folder("/root", when=when)
        self.assertEqual(path, "/root/spravy/02.09.2026/prohin_14-30")

    def test_run_id_names_the_subfolder(self):
        path = E.run_folder("/root", "run-20260902-143000-ab12cd")
        self.assertTrue(path.endswith("prohin_ab12cd"), path)

    def test_two_runs_the_same_day_do_not_mix(self):
        """Два прогони — дві теки; документи качаються ОДИН раз."""
        st1 = self.download(make_pdf_copies=False)
        calls = []
        net2 = types.SimpleNamespace(
            make_session=lambda: object(),
            get_bytes=lambda s, u, tries=3: (calls.append(u) or
                                             self.bytes.get(u, b"%PDF-1.4 x")))
        second = E.run_folder(self.tmp, "run-second-xyz789")
        st2 = E.download_case_documents(self.db, self.tmp, legacy=net2,
                                        case_dir=second, make_pdf_copies=False)
        # Теку прогону створює покажчик: порожніх тек «про всяк випадок»
        # движок не робить.
        for folder in (self.case, second):
            E.build_case_cards(self.db, self.tmp, legacy=L, case_dir=folder)
        day = os.path.dirname(self.case)
        runs = sorted(os.listdir(day))
        self.assertEqual(len(runs), 2, f"прогони змішались: {runs}")
        self.assertGreater(st1["завантажено"], 0)
        self.assertEqual(calls, [], "другий прогін качав ті самі документи знову")
        self.assertEqual(st2["завантажено"], 0)
        self.assertEqual(st2["вже_були"], st1["завантажено"])
        self.assertTrue(os.path.isfile(os.path.join(self.folder,
                                                    E.MANIFEST_NAME)))

    def test_the_run_folder_gets_an_index(self):
        """Тека за датою тримає покажчик, а не копії документів."""
        self.download(make_pdf_copies=False)
        E.build_case_cards(self.db, self.tmp, legacy=L, case_dir=self.case)
        index = os.path.join(self.case, E.RUN_INDEX_NAME)
        self.assertTrue(os.path.isfile(index), sorted(os.listdir(self.case)))
        import docx
        text = "\n".join(c.text for t in docx.Document(index).tables
                         for r in t.rows for c in r.cells)
        self.assertIn("UA-2026-08-07-008904-a", text)
        # У теці прогону не має бути жодного завантаженого документа
        self.assertEqual([n for n in os.listdir(self.case)
                          if n != E.RUN_INDEX_NAME], [])

    def test_case_folders_reports_across_runs(self):
        import contextlib
        self.download(make_pdf_copies=False)
        E.build_case_cards(self.db, self.tmp, legacy=L, case_dir=self.case)
        buf = io.StringIO()
        with contextlib.redirect_stdout(buf):
            E.case_folders(self.tmp)
        out = buf.getvalue()
        self.assertIn("прогін:", out)
        self.assertIn("UA-2026-08-07-008904-a", out)
        self.assertIn("КАРТКА:", out)


if __name__ == "__main__":
    unittest.main(verbosity=2)
